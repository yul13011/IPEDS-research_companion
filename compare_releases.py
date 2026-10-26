from pathlib import Path
import csv,json,re,collections,zipfile,xml.etree.ElementTree as ET,difflib
import openpyxl

BASE=Path(__file__).parent
ROOT=BASE.parent/'IPEDS Data'
OUT=BASE/'release-comparison';OUT.mkdir(exist_ok=True)
def text(v):return '' if v is None else str(v)
def clean(v):return re.sub(r'\s+',' ',text(v).replace('_x000D_',' ').replace('\u00a0',' ')).strip()
def family(name,year):
    s=name.upper();start=int(year[6:10]);a=str(start)[-2:];b=str(start+1)[-2:]
    s=s.replace(str(start)+str(start+1),'{CYCLE}').replace(str(start)+b,'{CYCLE}')
    for x in [str(start-1)[-2:]+a,a+b]:s=re.sub(r'^(SFA|SFAV|F)'+x+r'(?=_|$)',r'\1{PERIOD}',s)
    s=s.replace(str(start),'{YEAR}')
    s=re.sub(r'(?<=_)'+a+r'(?=$|_)','{YY}',s)
    s=re.sub(r'(?<=[A-Z])'+a+r'$','{YY}',s)
    # Metadata tables use Tables15, valuesets15 etc.
    return s
def norm(v):
    s=clean(v).lower()
    s=re.sub(r'\b(?:19|20)\d{2}\b','<year>',s)
    s=re.sub(r'\b((?:19|20)\d{2})?\s*[-–]\s*\d{2}\b','<period>',s)
    s=re.sub(r'((?:gr|ef|ic|sfa|f|c)\d{2,8})',lambda m:re.sub(r'\d+','<year>',m[0]),s)
    return s
def writecsv(name,rows,fields):
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)

schemas=collections.defaultdict(lambda:collections.defaultdict(dict))
with (OUT/'database_schema.csv').open(encoding='utf-8',newline='') as f:
    for r in csv.DictReader(f):schemas[r['release']][family(r['table'],r['release'])][r['column'].upper()]=r
releases=[];notes=[];metadata=[];sheets=[]
for folder in sorted(ROOT.iterdir()):
    if not folder.is_dir():continue
    release=folder.name;book=next(folder.glob('*.xlsx'));w=openpyxl.load_workbook(book,read_only=True,data_only=True)
    records={};intro=[]
    for s in w:
        rows=list(s.values);sheets.append({'release':release,'sheet':s.title,'rows':s.max_row,'columns':s.max_column})
        kind=next((k for k in ['vartable','valuesets','tables'] if s.title.lower().startswith(k)),None)
        if kind:
            header=[clean(x).lower() for x in rows[0]]
            records[kind]=[(n,dict(zip(header,r))) for n,r in enumerate(rows[1:],2) if any(x is not None for x in r)]
        else:
            (OUT/(release+'_'+re.sub(r'[^a-zA-Z0-9_-]','_',s.title)+'.json')).write_text(json.dumps(rows,default=str,ensure_ascii=False),encoding='utf-8')
            intro.extend(clean(x) for r in rows for x in r if x is not None)
    (OUT/(release+'_workbook_notes.txt')).write_text('\n'.join(intro),encoding='utf-8')
    for doc in folder.glob('*.docx'):
        with zipfile.ZipFile(doc) as z:
            root=ET.fromstring(z.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            paragraphs=[''.join(t.text or '' for t in p.findall('.//w:t',ns)) for p in root.findall('.//w:p',ns)]
        notes.append({'release':release,'file':doc.name,'text':'\n'.join(paragraphs)})
        (OUT/(release+'_readme.txt')).write_text('\n'.join(paragraphs),encoding='utf-8')
    tables={};variables={};codes={}
    for n,d in records.get('tables',[]):
        t=clean(d.get('tablename'))
        if t:tables[family(t,release)]={'row':n,**{k:clean(v) for k,v in d.items() if k}}
    for n,d in records.get('vartable',[]):
        t=clean(d.get('tablename'));v=clean(d.get('varname')).upper()
        if t and v:variables[(family(t,release),v)]={'row':n,**{k:clean(x) for k,x in d.items() if k}}
    for n,d in records.get('valuesets',[]):
        t=clean(d.get('tablename'));v=clean(d.get('varname')).upper();code=clean(d.get('codevalue'))
        if t and v:codes[(family(t,release),v,code)]={'row':n,**{k:clean(x) for k,x in d.items() if k}}
    releases.append({'release':release,'book':str(book),'tables':tables,'variables':variables,'codes':codes})
    print(release,len(tables),len(variables),len(codes),flush=True)

writecsv('sheet_inventory.csv',sheets,['release','sheet','rows','columns'])
(OUT/'readmes.json').write_text(json.dumps(notes,indent=2),encoding='utf-8')
changes=[];summary=[];mismatches=[]
for i,r in enumerate(releases):
    year=r['release'];actual=schemas[year]
    for fam,t in r['tables'].items():
        if fam not in actual:mismatches.append({'release':year,'table':t['tablename'],'variable':'','issue':'Documented table absent from database','source_row':t['row']})
    for (fam,v),d in r['variables'].items():
        if fam in actual and v not in actual[fam]:mismatches.append({'release':year,'table':d['tablename'],'variable':v,'issue':'Documented variable absent from database','source_row':d['row']})
    for fam,cols in actual.items():
        if fam not in r['tables']:continue
        for v,d in cols.items():
            if (fam,v) not in r['variables']:mismatches.append({'release':year,'table':d['table'],'variable':v,'issue':'Database column absent from variable dictionary','source_row':''})
    if i==0:continue
    for mode,b in [('baseline',releases[0]),('previous',releases[i-1])]:
        start=len(changes)
        def add(kind,key,field,old,new,oldrow='',newrow=''):
            changes.append({'comparison':mode,'from_release':b['release'],'to_release':year,'kind':kind,'table_family':key[0] if isinstance(key,tuple) else key,'variable':key[1] if isinstance(key,tuple) and len(key)>1 else '','code':key[2] if isinstance(key,tuple) and len(key)>2 else '','field':field,'old':text(old),'new':text(new),'old_source_row':oldrow,'new_source_row':newrow})
        for kind,fields in [('tables',['tabletitle','description','yearcoverage']),('variables',['vartitle','longdescription','datatype','format','fieldwidth','multirecord','imputationvar']),('codes',['valuelabel'])]:
            old=b[kind];new=r[kind]
            for k in old.keys()-new.keys():add(kind+'_removed',k,'',old[k].get('vartitle',old[k].get('tabletitle',old[k].get('valuelabel'))),'',old[k]['row'])
            for k in new.keys()-old.keys():add(kind+'_added',k,'','',new[k].get('vartitle',new[k].get('tabletitle',new[k].get('valuelabel'))),'',new[k]['row'])
            for k in old.keys()&new.keys():
                for field in fields:
                    a=old[k].get(field,'');z=new[k].get(field,'')
                    if clean(a)==clean(z):continue
                    category='period_or_wording' if norm(a)==norm(z) else 'definition_or_structure'
                    add(kind+'_'+category,k,field,a,z,old[k]['row'],new[k]['row'])
        old=schemas[b['release']]
        for fam in old.keys()&actual.keys():
            for v in old[fam].keys()-actual[fam].keys():add('db_column_removed',(fam,v),'type',old[fam][v]['type'],'')
            for v in actual[fam].keys()-old[fam].keys():add('db_column_added',(fam,v),'type','',actual[fam][v]['type'])
            for v in old[fam].keys()&actual[fam].keys():
                if old[fam][v]['type']!=actual[fam][v]['type']:add('db_type_changed',(fam,v),'type',old[fam][v]['type'],actual[fam][v]['type'])
        cnt=collections.Counter(x['kind'] for x in changes[start:]);summary.append({'from_release':b['release'],'to_release':year,'comparison':mode,'counts':dict(cnt)})
writecsv('change_register.csv',changes,['comparison','from_release','to_release','kind','table_family','variable','code','field','old','new','old_source_row','new_source_row'])
writecsv('metadata_database_mismatches.csv',mismatches,['release','table','variable','issue','source_row'])
(OUT/'comparison_summary.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
catalog=[]
for r in releases:
    yr=r['release'];ts=[]
    for fam,t in r['tables'].items():
        cols=schemas[yr].get(fam,{})
        ts.append({'family':fam,'name':t['tablename'],'title':t.get('tabletitle'),'description':t.get('description'),'period':t.get('yearcoverage'),'source_row':t['row'],'columns':len(cols),'rows':next(iter(cols.values()))['source_rows'] if cols else None})
    catalog.append({'release':yr,'tables':ts,'variable_entries':len(r['variables']),'code_entries':len(r['codes']),'database_tables':len(schemas[yr])})
(OUT/'catalog_snapshot.json').write_text(json.dumps(catalog,indent=2),encoding='utf-8')
print('CHANGES',len(changes),'MISMATCHES',len(mismatches))

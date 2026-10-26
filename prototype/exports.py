"""Exports contain full results plus an audit trail; never a preview-only download."""
from io import BytesIO
import json
import pandas as pd
from .units import display_result


def dataframe(result):
    return pd.DataFrame(display_result(result)['rows'])


def _safe(value):
    if isinstance(value,(dict,list)):value=json.dumps(value,ensure_ascii=False)
    # Prevent spreadsheet formula interpretation in text originating in source files.
    return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value


def csv_bytes(result):
    return dataframe(result).map(_safe).to_csv(index=False).encode('utf-8-sig')


def audit_bytes(result):
    rows=[]
    for row in result['rows']:
        raw=dict(row)
        if 'source_unit' in raw:
            raw['unit']=raw.pop('source_unit')
            raw.pop('currency',None);raw.pop('currency_scale',None)
        rows.append(raw)
    return json.dumps({**result,'rows':rows,'display_units':display_result(result)['display_units']},indent=2,ensure_ascii=False).encode('utf-8')


def excel_bytes(result):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.title='Data'
    frame=dataframe(result);ws.append(list(frame.columns))
    for row in frame.to_dict('records'):ws.append([_safe(row.get(c)) for c in frame.columns])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(55,max(16,max(len(str(c.value or '')) for c in col)+2))
    meta=wb.create_sheet('Methodology');meta.append(['Item','Value'])
    for k in ['result_id','created_at','registry_sha256','review_status','plan','coverage','definitions','warnings']:
        value=result[k];meta.append([k,_safe(json.dumps(value,ensure_ascii=False) if not isinstance(value,str) else value)])
    for k in ['catalog_plan','semantic_validation','calculation_validation','latency','technical_notes']:
        if k in result:meta.append([k,json.dumps(result[k],ensure_ascii=False)])
    meta.column_dimensions['A'].width=25;meta.column_dimensions['B'].width=100
    meta.append(['display_units',json.dumps(display_result(result)['display_units'],ensure_ascii=False)])
    sources=wb.create_sheet('Sources');fields=['metric','year','unitid','database','database_size','database_mtime_ns','table','columns','filters','formula','sql','parameters']
    sources.append(fields)
    for r in result['sources']:sources.append([_safe(json.dumps(r[k],ensure_ascii=False) if isinstance(r[k],(dict,list)) else r[k]) for k in fields])
    definitions=wb.create_sheet('Codebook evidence');definitions.append(['metric','year','record_id','variable','title','definition','source_file','sheet','row','documented_table','physical_table','issues','code_filters'])
    seen=set()
    for source in result['sources']:
        for doc in source['codebook_evidence']:
            key=(source['metric'],source['year'],doc['record_id'])
            if key in seen:continue
            seen.add(key)
            definitions.append([_safe(v) for v in [source['metric'],source['year'],doc['record_id'],doc['variable_name'],doc['title'],doc['long_description'],doc['source_file'],doc['sheet_name'],doc['source_row'],doc['documented_table'],doc['physical_table'],json.dumps(doc['issues']),json.dumps(source['code_filters'])]])
    if result.get('population_evidence'):
        scope=wb.create_sheet('Aggregate scope');scope.append(['Source year','Scope','Directory count','Filter codes','Directory SQL','Parameters'])
        units=wb.create_sheet('Population reporting units');units.append(['Source year','UNITID'])
        for p in result['population_evidence'].values():
            scope.append([p['year'],p['label'],p['directory_count'],json.dumps(p['filters']),p['directory_sql'],json.dumps(p['parameters'])])
            for uid in p['unitids']:units.append([p['year'],uid])
    out=BytesIO();wb.save(out);return out.getvalue()

"""Real-source execution checks; no live model calls."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sys
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[2]
sys.path[:0]=[str(ROOT/'.packages'),str(ROOT)]
from prototype.agent import ask
from prototype.catalog import annual_path, readonly
from prototype.catalog_planner import enrollment_race_plan, proposal_answer
from prototype.catalog_execution import execute_catalog, numeric
from prototype.exports import csv_bytes, excel_bytes, audit_bytes
from analysis_catalog.catalog import connection


class CatalogExecutionTests(unittest.TestCase):
    def test_race_break_down_wording_uses_same_source_plan(self):
        questions=["show Yale's undergraduate population's race/ethnicity break down.",
                   "Show Yale’s undergraduate population’s race/ethnicity break-down."]
        baseline=ask("Show Yale's undergrad population by race and ethnicity")
        with patch('prototype.catalog_planner.model_catalog_plan',side_effect=AssertionError('Equivalent wording must not require a model-generated grouped plan')):
            for question in questions:
                with self.subTest(question=question):
                    answer=ask(question,use_model=True)
                    self.assertEqual(answer['intent'],'analyze')
                    self.assertEqual(answer['catalog_plan'],baseline['catalog_plan'])
                    self.assertEqual(answer['result']['rows'],baseline['result']['rows'])
        self.assertIsNone(enrollment_race_plan("Show Yale's undergraduate population race/ethnicity break down by income"))

    def test_race_count_matches_independent_source_query(self):
        answer=ask("Show Yale's undergrad population by race and ethnicity")
        self.assertEqual(answer['intent'],'analyze')
        result=answer['result'];self.assertTrue(result['semantic_validation']['execution_approved'])
        with readonly(annual_path(2024)) as db:
            row=db.execute('SELECT EFASIAT,EFTOTLT FROM EF2024A WHERE UNITID=? AND EFALEVEL=?',(130794,2)).fetchone()
        self.assertEqual(row['EFTOTLT'],6814)
        with connection() as db:
            code=db.execute('SELECT value_label FROM source.code_values WHERE record_id=?',(answer['catalog_plan']['filters'][0]['code_record_ids'][0],)).fetchone()[0]
        self.assertEqual(code,'All students, Undergraduate total')
        self.assertEqual(next(r['value'] for r in result['rows'] if r['measure']=='Asian total'),row[0])
        self.assertEqual(next(r['value'] for r in result['rows'] if r['measure']=='Grand total'),row[1])
        self.assertGreaterEqual(answer['latency']['answer_ready_seconds'],answer['latency']['query_seconds'])

    def test_percent_exact_denominator_and_missing_values(self):
        answer=ask("Show Yale's undergrad population percentages by race and ethnicity in 2024")
        rows=answer['result']['rows'];self.assertEqual(len(rows),9)
        self.assertAlmostEqual(sum(r['value'] for r in rows),100)
        asian=next(r for r in rows if 'EFASIAT' in r['measure'])
        self.assertAlmostEqual(asian['value'],100*asian['numerator']/asian['denominator'])
        for v in [None,'',-1,-2,-3,-9,'NaN','Infinity','suppressed']:self.assertIsNone(numeric(v))
        self.assertEqual(numeric(0),0);self.assertEqual(numeric('1.25'),1.25)

    def test_continuing_subgroup_cannot_answer_all_undergraduates(self):
        question="Show Yale's undergrad population by race and ethnicity"
        answer=enrollment_race_plan(question)
        with connection() as db:
            code=db.execute("SELECT record_id FROM source.code_values WHERE release_id='IPEDS_2024-25_Provisional' AND table_name='EF2024A' AND variable_name='EFALEVEL' AND code_value='20'").fetchone()[0]
        answer['catalog_plan']['filters'][0].update(values=['20'],code_record_ids=[code])
        result=execute_catalog(answer,question)
        self.assertNotIn('result',result)
        self.assertIn('does not represent all undergraduates',result['message'])

    def test_unselected_population_never_executes(self):
        answer=enrollment_race_plan("Show Yale's undergrad population by race and ethnicity")
        answer['catalog_plan']['filters']=[]
        with patch('prototype.catalog_execution.readonly',side_effect=AssertionError('Blocked plans cannot query')):
            result=execute_catalog(answer,'Yale undergraduate population')
        self.assertNotIn('result',result);self.assertIn('EFALEVEL',result['message'])

    def test_unverified_arithmetic_and_dimensions_blocked(self):
        base=enrollment_race_plan("Show Yale's undergrad population percentages by race and ethnicity")
        for mutation in ['calculation','dimension']:
            answer=deepcopy(base)
            if mutation=='calculation':answer['catalog_plan']['calculations'][0]['operator']='sum'
            else:answer['catalog_plan']['selections'][0]['role']='dimension'
            self.assertNotIn('result',execute_catalog(answer,'Yale undergraduate population'))

    def test_unfamiliar_library_field_without_registry_entry(self):
        with connection() as db:
            rid=db.execute("SELECT record_id FROM definitions WHERE release_id='IPEDS_2024-25_Provisional' AND table_name='AL2024' AND variable_name='LPBOOKS'").fetchone()[0]
        p={'version':1,'institutions':['Yale'],'source_years':[2024],
           'requested_time':{'basis':'fiscal','description':'Fiscal year 2024'},'selections':[{'record_id':rid,'role':'measure'}],
           'filters':[],'group_by_record_ids':[],'calculations':[],'joins':[],'output':'table','assumptions':[],'clarification_questions':[]}
        answer=execute_catalog(proposal_answer(p,'Yale physical library books fiscal 2024'),'Yale physical library books fiscal 2024')
        self.assertIn('result',answer,answer['message'])
        with readonly(annual_path(2024)) as db:expected=db.execute('SELECT LPBOOKS FROM AL2024 WHERE UNITID=130794').fetchone()[0]
        self.assertEqual(answer['result']['rows'][0]['value'],expected)
        result=answer['result']
        self.assertIn(b'Number of physical books',csv_bytes(result));self.assertIn(b'semantic_validation',audit_bytes(result))
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(excel_bytes(result)),read_only=True)
        self.assertIn('Codebook evidence',book.sheetnames)
        self.assertEqual(book['Data'].max_row,2)

    def test_readonly_connection_rejects_write(self):
        import sqlite3
        with readonly(annual_path(2024)) as db:
            with self.assertRaises(sqlite3.OperationalError):db.execute('CREATE TABLE should_never_exist (id INTEGER)')

    def test_codebook_and_physical_table_case_differences(self):
        with connection() as db:
            rid=db.execute("SELECT record_id FROM definitions WHERE release_id='IPEDS_2024-25_Provisional' AND table_name='COST1_2024' COLLATE NOCASE AND variable_name='TUITION3'").fetchone()[0]
        p={'version':1,'institutions':['Yale'],'source_years':[2024],
           'requested_time':{'basis':'academic','description':'Academic year 2024-25'},'selections':[{'record_id':rid,'role':'measure'}],
           'filters':[],'group_by_record_ids':[],'calculations':[],'joins':[],'output':'table','assumptions':[],'clarification_questions':[]}
        answer=execute_catalog(proposal_answer(p,'Yale out-of-state undergraduate tuition 2024-25'),'Yale out-of-state undergraduate tuition 2024-25')
        self.assertIn('result',answer,answer['message'])
        with readonly(annual_path(2024)) as db:expected=db.execute('SELECT TUITION3 FROM Cost1_2024 WHERE UNITID=130794').fetchone()[0]
        self.assertEqual(answer['result']['rows'][0]['value'],expected)

    def test_unknown_field_and_sql_text_rejected(self):
        answer=enrollment_race_plan("Show Yale's undergrad population by race and ethnicity")
        answer['catalog_plan']['selections'][0]['record_id']=999999999
        with self.assertRaises(ValueError):execute_catalog(answer,'Yale population')


if __name__=='__main__':unittest.main()

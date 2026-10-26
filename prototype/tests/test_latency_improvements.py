"""Latency changes must preserve evidence, identity checks, and file contents."""
from pathlib import Path
from copy import deepcopy
import json,sys,time,unittest
from unittest.mock import Mock,patch
ROOT=Path(__file__).resolve().parents[2];sys.path[:0]=[str(ROOT/'.packages'),str(ROOT)]

class LatencyTests(unittest.TestCase):
    def test_identity_extraction_reuses_claude_interpretation_and_confirms_directory(self):
        from prototype.intent_router import INTERPRETED
        from prototype.institution_confirmation import mentions,resolve
        token=INTERPRETED.set({'institution_mentions':['Yale University','Harvard']})
        try:
            with patch('anthropic.Anthropic',side_effect=AssertionError('Duplicate model call')):
                names,usage=mentions('Compare Yale University and Harvard in 2024',True,'test')
            self.assertEqual(set(names),{'Yale University','Harvard'});self.assertEqual(usage['calls'],0)
            self.assertEqual(resolve('Yale University',{}),130794)
            self.assertIsNone(resolve('Michigan',{}))
        finally:INTERPRETED.reset(token)

    def test_http_clients_reuse_connections_and_isolate_credentials_and_factories(self):
        from prototype.api_client import get_client,close_clients
        close_clients()
        with patch('anthropic.Anthropic') as factory,patch('prototype.agent.api_key',return_value='test-key-one'):
            a=get_client();b=get_client();self.assertIs(a,b);factory.assert_called_once()
            with patch('prototype.agent.api_key',return_value='test-key-two'):get_client()
            self.assertEqual(factory.call_count,2)
        with patch('anthropic.Anthropic') as other,patch('prototype.agent.api_key',return_value='test-key-one'):
            get_client();other.assert_called_once()
        close_clients()

    def test_compact_writer_retains_every_passage_and_valid_source_citation(self):
        from prototype.discovery import overview
        from prototype.discovery_narrative import evidence_package,compact_package,restore_citations,validate
        from prototype.tests.test_discovery_narrative import valid_draft
        for topic in ['faculty','graduation rates']:
            with self.subTest(topic=topic):
                original,_=evidence_package(overview('What data do you have on '+topic+'?'),topic)
                before=deepcopy(original);compact,aliases=compact_package(original)
                self.assertEqual({p['text'] for r in original['source_records'] for p in r['passages']},{p['text'] for p in compact['source_passages']})
                self.assertEqual(compact['topics'],original['topics'])
                self.assertEqual(compact['required_qualifications'],original['required_qualifications'])
                self.assertLess(len(json.dumps(compact)),len(json.dumps(original))*(.65 if topic=='faculty' else .75))
                draft=valid_draft(original)
                original_text={p['quote_id']:p['text'] for r in original['source_records'] for p in r['passages']}
                for b in draft['blocks']:
                    for p in b['points']:
                        for c in p['citations']:
                            c['quote_id']=next(p['quote_id'] for p in compact['source_passages'] if p['text']==original_text[c['quote_id']] and c['topic_id'] in p['topic_ids'])
                restored=restore_citations(draft,aliases);validate(restored,original)
                self.assertEqual(original,before)
                with self.assertRaises(ValueError):restore_citations({'blocks':[{'points':[{'citations':[{'topic_id':'Missing','quote_id':'C1'}]}]}]},aliases)

    def test_fast_series_names_equal_actual_rendered_domains(self):
        from prototype.analytics import Plan,execute
        from prototype.responses import finalize_answer
        from prototype.charts import series_names_for_context
        from prototype.chart_format import series_names
        from prototype.workflows import execute_example
        results=[execute(Plan(['admission_rate'],['Ivy League'],[2024])),
                 execute(Plan(['total_enrollment'],['Yale','Harvard'],list(range(2015,2025)))),
                 execute_example('faculty-salaries')['result']]
        for result in results:
            for chart in ['line','bar']:
                for separate in [True,False]:
                    p={**finalize_answer({'result':result})['presentation'],'chart_type':chart,'separate_panels':separate}
                    self.assertEqual(series_names_for_context(result,p),series_names(result,p))

    def test_deferred_downloads_generate_identical_files_only_on_request(self):
        from streamlit.testing.v1 import AppTest
        from prototype.analytics import Plan,execute
        from prototype.responses import finalize_answer
        from prototype import exports
        from prototype.components import render_result
        from io import BytesIO
        from openpyxl import load_workbook
        answer=finalize_answer({'result':execute(Plan(['admission_rate'],['Ivy League'],[2024]))})
        captured=[]
        def button(label,data,**kw):captured.append((label,data,kw));return False
        app=AppTest.from_file(str(ROOT/'prototype/app.py'),default_timeout=60).run()
        app.session_state['use_ai']=False;app.session_state['main_page']='Ask & Analyze'
        app.session_state['history']=[{'question':'Admission rates','answer':answer}]
        with patch('streamlit.download_button',side_effect=button),patch('prototype.components.excel_bytes',wraps=exports.excel_bytes) as excel:
            app.run();self.assertFalse(app.exception);excel.assert_not_called()
            callbacks={label:data for label,data,_ in captured}
            self.assertTrue(all(callable(data) for _,data,_ in captured))
            self.assertTrue(all(kw.get('on_click')=='ignore' for _,_,kw in captured))
            sheet=load_workbook(BytesIO(callbacks['Excel + sources']()))
            self.assertEqual(sheet['Data'].max_row,9);excel.assert_called_once()
            self.assertEqual(callbacks['CSV'](),exports.csv_bytes(answer['result']))
            self.assertEqual(json.loads(callbacks['Technical audit file (JSON)']()),json.loads(exports.audit_bytes(answer['result'])))

if __name__=='__main__':unittest.main()

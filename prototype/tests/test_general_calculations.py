"""Cross-survey calculation proofs, independent SQL checks and bounded repair."""
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch, MagicMock
import json
import sys
import unittest

ROOT=Path(__file__).resolve().parents[2]
sys.path[:0]=[str(ROOT/'.packages'),str(ROOT)]
from analysis_catalog.catalog import connection
from prototype.catalog import annual_path, readonly
from prototype.catalog_planner import proposal_answer, model_catalog_plan, read_definition
from prototype.catalog_execution import execute_catalog, prepare_catalog_proposal
from prototype.tests.test_catalog_planner import response, block


def plan_for(table,inputs,denominator=None,operator='percent',selectors=None,basis='reporting_period',year=2024):
    with connection() as db:
        release=db.execute('SELECT release_id FROM tables WHERE year=? AND table_name=? COLLATE NOCASE',(year,table)).fetchone()[0]
        definitions={r['variable_name']:dict(r) for r in db.execute('SELECT * FROM definitions WHERE release_id=? AND table_name=? COLLATE NOCASE',(release,table))}
        filters=[]
        for col,value in (selectors or {}).items():
            code=db.execute('SELECT record_id FROM source.code_values WHERE release_id=? AND table_name=? COLLATE NOCASE AND variable_name=? AND code_value=?',(release,table,col,value)).fetchone()[0]
            filters.append(dict(record_id=definitions[col]['record_id'],operator='eq',values=[value],code_record_ids=[code]))
    selections=[dict(record_id=definitions[c]['record_id'],role='measure') for c in inputs]
    if denominator:selections.append(dict(record_id=definitions[denominator]['record_id'],role='denominator'))
    return dict(version=1,institutions=['Yale'],source_years=[year],requested_time=dict(basis=basis,description='Exact source reporting period'),
        selections=selections,filters=filters,group_by_record_ids=[],calculations=[dict(label=operator+' of '+', '.join(inputs),operator=operator,
        input_record_ids=[definitions[c]['record_id'] for c in inputs],denominator_record_id=definitions[denominator]['record_id'] if denominator else None,
        group_by_record_ids=[],interpretation='Source-derived arithmetic')],joins=[],output='table',assumptions=[],clarification_questions=[])


class GeneralCalculationTests(unittest.TestCase):
    def run_plan(self,plan,question='Yale source category share'):
        return execute_catalog(proposal_answer(plan,question),question)

    def test_admissions_share_not_in_old_field_allowlist(self):
        p=plan_for('ADM2024',['APPLCNW'],'APPLCN',basis='fall')
        a=self.run_plan(p,'What percentage of Yale applicants were women in fall 2024?')
        self.assertTrue('result' in a,a['message'])
        with readonly(annual_path(2024)) as db:n,d=db.execute('SELECT APPLCNW,APPLCN FROM ADM2024 WHERE UNITID=130794').fetchone()
        self.assertAlmostEqual(a['result']['rows'][0]['value'],100*n/d)
        self.assertEqual(a['result']['calculation_validation'][0]['proof']['kind'],'category_containment')
        from prototype.exports import excel_bytes,csv_bytes,audit_bytes
        from openpyxl import load_workbook
        from io import BytesIO
        workbook=load_workbook(BytesIO(excel_bytes(a['result'])),read_only=True)
        self.assertEqual(workbook['Data'].max_row,2)
        self.assertIn(b'category_containment',csv_bytes(a['result']))
        self.assertIn(b'field_semantics',audit_bytes(a['result']))

    def test_completions_share_with_all_grain_selectors(self):
        p=plan_for('C2024_A',['CTOTALW'],'CTOTALT',selectors={'AWLEVEL':'5','CIPCODE':'99','MAJORNUM':'1'},basis='academic')
        a=self.run_plan(p,"Yale women as a percentage of bachelor's degrees awarded in 2023–24")
        self.assertTrue('result' in a,a['message'])
        with readonly(annual_path(2024)) as db:n,d=db.execute("SELECT CTOTALW,CTOTALT FROM C2024_A WHERE UNITID=130794 AND AWLEVEL=5 AND CIPCODE='99' AND MAJORNUM=1").fetchone()
        self.assertAlmostEqual(a['result']['rows'][0]['value'],100*n/d)
        missing=deepcopy(p);missing['filters']=missing['filters'][:1]
        self.assertNotIn('result',self.run_plan(missing))

    def test_unseen_institution_and_annual_definition_use_same_rules(self):
        from prototype.catalog import institutions
        p=plan_for('ADM2015',['APPLCNW'],'APPLCN',basis='fall',year=2015)
        p['institutions']=['Harvard']
        a=self.run_plan(p,'Share of Harvard applicants who were women in fall 2015')
        self.assertTrue('result' in a,a['message'])
        unitid=institutions(['Harvard'],2015)[0]
        with readonly(annual_path(2015)) as db:n,d=db.execute('SELECT APPLCNW,APPLCN FROM ADM2015 WHERE UNITID=?',(unitid,)).fetchone()
        self.assertAlmostEqual(a['result']['rows'][0]['value'],100*n/d)

    def test_unrequested_calculation_rejected_for_replanning(self):
        p=plan_for('ADM2024',['APPLCNW'],'APPLCN')
        a=self.run_plan(p,'Show Yale applicants by gender in fall 2024')
        self.assertNotIn('result',a)
        self.assertIn('does not ask for this derived calculation',a['message'])

    def test_catalog_descriptions_cover_definitions_without_blanket_approval(self):
        import sqlite3
        from contextlib import closing
        with connection() as db:expected=db.execute('SELECT COUNT(*) FROM definitions').fetchone()[0]
        with closing(sqlite3.connect((ROOT/'data/field_semantics.sqlite').as_uri()+'?mode=ro',uri=True)) as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM field_semantics').fetchone()[0],expected)
            descriptors=[json.loads(r[0]) for r in db.execute('SELECT descriptor_json FROM field_semantics')]
        self.assertTrue(any(d['unit']=='unknown' for d in descriptors))
        self.assertTrue(all(d['execution_approved'] is False for d in descriptors))

    def test_library_share_and_explicit_source_component(self):
        for numerator,denominator in [('LPCLLCT','LTCLLCT'),('LPBOOKS','LPCLLCT')]:
            a=self.run_plan(plan_for('AL2024',[numerator],denominator,basis='fiscal'))
            self.assertTrue('result' in a,a['message'])
            with readonly(annual_path(2024)) as db:n,d=db.execute(f'SELECT {numerator},{denominator} FROM AL2024 WHERE UNITID=130794').fetchone()
            self.assertAlmostEqual(a['result']['rows'][0]['value'],100*n/d)
        self.assertEqual(a['result']['calculation_validation'][0]['proof']['kind'],'explicit_source_decomposition')

    def test_general_sum_difference_and_sibling_ratio(self):
        for op,expected in [('sum',lambda m,w:m+w),('difference',lambda m,w:m-w),('ratio',lambda m,w:m/w)]:
            p=plan_for('ADM2024',['APPLCNM','APPLCNW'] if op!='ratio' else ['APPLCNM'],'APPLCNW' if op=='ratio' else None,operator=op)
            a=self.run_plan(p,'Yale applicant category '+op)
            self.assertTrue('result' in a,a['message'])
            with readonly(annual_path(2024)) as db:m,w=db.execute('SELECT APPLCNM,APPLCNW FROM ADM2024 WHERE UNITID=130794').fetchone()
            self.assertAlmostEqual(a['result']['rows'][0]['value'],expected(m,w))
        # These are source Men/Women counts, not an assertion about unknown/another gender.

    def test_false_relationships_and_overlapping_totals_rejected(self):
        cases=[plan_for('ADM2024',['APPLCNW'],'ADMSSN'),
               plan_for('ADM2024',['APPLCNW'],'APPLCNM'),
               plan_for('AL2024',['LPBOOKS'],'LEXPTOT'),
               plan_for('ADM2024',['APPLCNM','APPLCN'],operator='sum'),
               plan_for('ADM2024',['APPLCNM','APPLCNW'],operator='mean')]
        for p in cases:
            with self.subTest(p=p['calculations'][0]['label']):self.assertNotIn('result',self.run_plan(p))

    def test_source_period_and_cross_table_evidence_rejected(self):
        p=plan_for('AL2024',['LPCLLCT'],'LTCLLCT')
        # Physical presence plus a model claim cannot authorize cross-table ratios.
        with connection() as db:
            rid=db.execute("SELECT record_id FROM definitions WHERE release_id='IPEDS_2024-25_Provisional' AND table_name='ADM2024' AND variable_name='APPLCN'").fetchone()[0]
        p['selections'][-1]['record_id']=rid;p['calculations'][0]['denominator_record_id']=rid
        self.assertNotIn('result',self.run_plan(p))
        p=plan_for('AL2024',['LPCLLCT'],'LTCLLCT',basis='fall')
        self.assertNotIn('result',self.run_plan(p))

    def test_multi_category_normalization_and_constant_dimension(self):
        p=plan_for('ADM2024',['APPLCNM','APPLCNW'],'APPLCN')
        a=self.run_plan(p,'Yale applicants by gender as percentages')
        self.assertTrue('result' in a,a['message']);self.assertEqual(len(a['result']['rows']),2)
        self.assertTrue(a['catalog_proposal']['normalization'])
        from prototype.catalog_planner import enrollment_gender_plan
        p=enrollment_gender_plan('Yale first-year undergrads by gender')['catalog_plan']
        selector=p['filters'][0]['record_id']
        p['selections'].append({'record_id':selector,'role':'dimension'});p['group_by_record_ids']=[selector]
        a=self.run_plan(p,'Yale first-year undergrads by gender')
        self.assertTrue('result' in a,a['message']);self.assertFalse(a['catalog_plan']['group_by_record_ids'])

    def test_missing_zero_and_bounds_without_writing_source(self):
        # Arithmetic consumes only revalidated source values. Inject a read-only
        # query cursor to exercise missing/zero/bounds without changing data.
        from contextlib import contextmanager
        p=plan_for('ADM2024',['APPLCNW'],'APPLCN')
        @contextmanager
        def fake_connection(path):
            class DB:
                def execute(self,sql,params):
                    if 'INSTNM' in sql:return SimpleNamespace(fetchone=lambda:['Yale University'])
                    return SimpleNamespace(fetchmany=lambda n:[{'APPLCNW':state[0],'APPLCN':state[1]}])
            yield DB()
        with patch('prototype.catalog_execution.readonly',fake_connection):
            for state in [(None,10),(2,0),(2,None)]:
                a=self.run_plan(p);self.assertIsNone(a['result']['rows'][0]['value'])
            state=(20,10)
            with self.assertRaisesRegex(ValueError,'exceeds'):self.run_plan(p)

    def test_planner_retrieves_and_repairs_semantic_error(self):
        bad=plan_for('ADM2024',['APPLCNW'],'ADMSSN')
        good=plan_for('ADM2024',['APPLCNW'],'APPLCN')
        with patch('anthropic.Anthropic') as cls:
            client=MagicMock();cls.return_value=client
            client.messages.create.side_effect=[response(block('submit_catalog_plan',bad,'bad')),
                response(block('read_definition',{'record_id':good['selections'][-1]['record_id'],'offset':0,'limit':4000},'inspect')),
                response(block('submit_catalog_plan',good,'good'))]
            a=model_catalog_plan('Yale applicants by gender as percentages')
        self.assertTrue(a['preflight_validation']['execution_eligible'])
        self.assertEqual(len(a['validation_attempts']),1)
        self.assertEqual(a['model_usage']['calls'],3)
        self.assertEqual(a['model_usage']['semantic_repair_attempts'],1)
        self.assertIn('analytical_description',read_definition(good['selections'][-1]['record_id']))

    def test_failed_repair_is_bounded_and_never_queries(self):
        bad=plan_for('ADM2024',['APPLCNW'],'ADMSSN')
        with patch('anthropic.Anthropic') as cls,patch('prototype.catalog_execution.readonly',side_effect=AssertionError('No statistics during repair')):
            cls.return_value.messages.create.side_effect=[response(block('submit_catalog_plan',bad,str(i))) for i in range(3)]
            a=model_catalog_plan('Yale applicants by gender as percentages')
        self.assertFalse(a['preflight_validation']['execution_eligible'])
        self.assertEqual(len(a['validation_attempts']),2)
        self.assertNotIn('result',a)


if __name__=='__main__':unittest.main()

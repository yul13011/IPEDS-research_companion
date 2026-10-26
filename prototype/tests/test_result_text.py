from pathlib import Path
from copy import deepcopy
from io import BytesIO
import sys,unittest
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[2]
sys.path[:0]=[str(ROOT/'.packages'),str(ROOT)]
from prototype.tests.test_general_calculations import plan_for
from prototype.catalog_planner import proposal_answer
from prototype.catalog_execution import execute_catalog
from prototype.result_text import visible_warnings,definition_summary,calculation_summary,evidence_summary,source_summary

class ResultTextTests(unittest.TestCase):
    def test_ten_year_sources_show_exact_table_metadata_once_per_release(self):
        from prototype.agent import ask
        from prototype.catalog import readonly,DATA
        r=ask('percentage of Yale undergraduates by gender in the past 10 years')['result']
        before=deepcopy(r);summary=source_summary(r)
        self.assertEqual(len(summary['sources']),10)
        self.assertEqual(len(summary['source_groups']),1)
        self.assertIn('Race/ethnicity, gender',summary['source_groups'][0]['title'])
        self.assertFalse(summary['source_notes'])
        with readonly(DATA/'metadata_catalog.sqlite') as db:
            for year,row in zip(range(2015,2025),summary['sources']):
                table=f'EF{year}A'
                doc=db.execute('SELECT survey,year_coverage FROM table_docs WHERE table_name=?',(table,)).fetchone()
                self.assertEqual(row['Survey'],doc['survey'])
                self.assertEqual(row['Year coverage'],doc['year_coverage'])
                self.assertEqual(row['Table'],table)
                self.assertIn(str(year),row['IPEDS release'])
        self.assertIn('Provisional',summary['sources'][-1]['IPEDS release'])
        self.assertNotIn('.xlsx',str(summary))
        self.assertEqual(r,before)

    def test_reviewed_and_joined_results_keep_distinct_executed_tables(self):
        from prototype.analytics import execute as reviewed,Plan
        from prototype.tests.test_query_engine import source,plan,join,expr
        from prototype.query_engine import execute,field_ref
        r=reviewed(Plan(['applicants','undergraduate_enrollment'],['Yale'],[2024]))
        self.assertEqual({row['Table'] for row in source_summary(r)['sources']},{s['table'] for s in r['sources']})
        a,da=source('ADM2024',sid='a',fields=['APPLCN'])
        e,de=source('EF2024A',sid='e',fields=['EFTOTLT'],filters={'EFALEVEL':'4'})
        p=plan([a,e],[('xratio','First-time students per applicant')],joins=[join('a','e')],expressions=[expr('ratio','divide',[field_ref(de['EFTOTLT'],'e'),field_ref(da['APPLCN'],'a')])])
        r=execute(p,'Yale first-time students per applicant in fall 2024')
        self.assertEqual({row['Table'] for row in source_summary(r)['sources']},{'ADM2024','EF2024A'})

    def test_missing_and_conflicting_metadata_are_explicit(self):
        r=self.result()
        with patch('prototype.result_text._table_metadata',return_value={}):summary=source_summary(r)
        self.assertEqual(summary['sources'][0]['Table'],'EF2024A')
        self.assertEqual(summary['sources'][0]['Year coverage'],'Not documented')
        self.assertTrue(any('unavailable' in note for note in summary['source_notes']))
        docs=[{'survey':'Fall Enrollment','year_coverage':'Fall 2024','title':'Enrollment'},
              {'survey':'Different survey','year_coverage':'Fall 2023','title':'Different title'}]
        with patch('prototype.result_text._table_metadata',return_value={('IPEDS_2024-25_Provisional','ef2024a'):docs}):summary=source_summary(r)
        self.assertEqual(summary['sources'][0]['Survey'],'Conflicting metadata')
        self.assertEqual(summary['sources'][0]['Year coverage'],'Conflicting metadata')
        self.assertTrue(any('Fall 2023' in note and 'Fall 2024' in note for note in summary['source_notes']))

    def result(self):
        p=plan_for('EF2024A',['EFTOTLW'],'EFTOTLT',selectors={'EFALEVEL':'2'},basis='fall')
        p['calculations'][0]['label']='Women as percentage of undergraduates'
        p['assumptions']=['INTERNAL same-row denominator proof and selected source codes']
        return execute_catalog(proposal_answer(p,'Yale undergraduate women share in Fall 2024'),'Yale undergraduate women share in Fall 2024')['result']

    def test_assumptions_are_audit_notes_and_real_caveats_survive(self):
        r=self.result();self.assertFalse(any('INTERNAL' in w for w in r['warnings']))
        self.assertTrue(any('INTERNAL' in w for w in r['technical_notes']))
        r['warnings']+=['Separate Yale College figures are not established.','Some counts do not reconcile.','Unexpected source conflict requiring review.']
        r['catalog_plan']['assumptions'].append('Unexpected source conflict requiring review.')
        shown=visible_warnings(r)
        self.assertTrue(any('Separate Yale College' in w for w in shown))
        self.assertTrue(any('do not reconcile' in w for w in shown))
        self.assertTrue(any('Unexpected source conflict' in w for w in shown))
        self.assertTrue(any('provisional' in w for w in shown))
        self.assertFalse(any('INTERNAL' in w for w in shown))

    def test_concise_definitions_formula_and_source_scope(self):
        r=self.result();before=deepcopy(r);metric=r['rows'][0]['metric']
        definitions=definition_summary(r,metric);formula=calculation_summary(r,metric);scope=evidence_summary(r)
        self.assertEqual(len(definitions),2)
        self.assertFalse(any('CREDIT -' in d or 'NOTE:' in d for d in definitions))
        self.assertIn('women',formula.lower());self.assertIn('× 100',formula)
        self.assertNotIn('EFTOTL',formula)
        self.assertIn('Fall 2024',scope['periods']);self.assertIn('All students, Undergraduate total',scope['population'])
        self.assertEqual(r,before)

    def test_full_originals_and_notes_remain_in_downloads(self):
        from prototype.exports import excel_bytes,audit_bytes
        from openpyxl import load_workbook
        r=self.result();raw=audit_bytes(r)
        self.assertIn(b'INTERNAL',raw);self.assertIn(b'EFTOTLW',raw);self.assertIn(b'CREDIT',raw)
        book=load_workbook(BytesIO(excel_bytes(r)),read_only=True)
        methodology={row[0]:row[1] for row in book['Methodology'].iter_rows(values_only=True)}
        self.assertIn('INTERNAL',methodology['technical_notes'])
        self.assertTrue(any('CREDIT' in str(row) for row in book['Codebook evidence'].iter_rows(values_only=True)))

    def test_version_two_explanations_stay_out_of_warnings(self):
        from prototype.tests.test_query_engine import source,plan,expr
        from prototype.query_engine import execute,field_ref
        s,d=source('ADM2024',fields=['APPLCN'])
        p=plan([s],[('xmean','Mean applicants')],expressions=[expr('mean','mean',[field_ref(d['APPLCN'],'s')])])
        p['assumptions']=['INTERNAL source binding proof']
        r=execute(p,'Mean Yale applicants')
        self.assertIn('INTERNAL source binding proof',r['technical_notes'])
        self.assertNotIn('INTERNAL source binding proof',visible_warnings(r))
        self.assertIn('mean',calculation_summary(r,'xmean'))

    def test_chat_shows_scope_warning_and_short_evidence(self):
        from streamlit.testing.v1 import AppTest
        from prototype.responses import finalize_answer
        r=self.result();r['warnings'].append('Separate Yale College figures are not established.')
        a=finalize_answer({'intent':'analyze','result':r})
        app=AppTest.from_file(str(ROOT/'prototype/app.py'),default_timeout=60).run()
        app.session_state['main_page']='Ask & Analyze';app.session_state['history']=[{'question':'Women percentage','answer':a}]
        app.run();self.assertFalse(app.exception)
        self.assertTrue(any('Separate Yale College' in w.value for w in app.warning))
        self.assertFalse(any('INTERNAL' in w.value for w in app.warning))
        shown=' '.join(x.value for x in app.markdown)+ ' '.join(x.value for x in app.caption)
        self.assertNotIn('Retrieved 1 of 1',shown);self.assertNotIn('varTable24',shown)
        self.assertNotIn('CREDIT -',shown);self.assertIn('Calculation:',shown)
        tables=[table.value for table in app.dataframe if 'IPEDS release' in table.value.columns]
        self.assertEqual(len(tables),1)
        self.assertEqual(tables[0].iloc[0]['Table'],'EF2024A')
        self.assertEqual(tables[0].iloc[0]['Survey'],'Fall Enrollment')
        self.assertEqual(tables[0].iloc[0]['Year coverage'],'Fall 2024')
        self.assertNotIn('Tablesdoc.xlsx',shown)
        self.assertEqual(len(app.get('popover')),2)

if __name__=='__main__':unittest.main()

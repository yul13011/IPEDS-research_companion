"""Integration checks against the real local sources and unchanged Yale golds."""
from io import BytesIO
from pathlib import Path
import json
import unittest
import sys

sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
sys.path.insert(0,str(Path(__file__).resolve().parents[2]/'.packages'))
from prototype.analytics import Plan, execute, measures
from prototype.agent import ask, local_plan, canonical_time
from prototype.catalog import BASE, DATA, readonly, institutions, search, definition_detail
from prototype.exports import csv_bytes, excel_bytes, audit_bytes


def fixture_metric(r):
    m=r['metric'].replace('Non Pell','Non-Pell')
    if r['family']=='admissions':return {'Total Applications':'applicants','Total Admits':'admits','Total Matrics':'matriculants','Admit Rate':'admission_rate','Yield Rate':'yield_rate'}[m]
    if r['family']=='enrollment':return {'Graduate/Professional':'graduate_enrollment','Undergraduate':'undergraduate_enrollment','Total':'total_enrollment'}[m]
    if r['family']=='sat':return m.replace('Verbal_','satvr').replace('Math_','satmt')
    if r['family']=='graduation':return {'Number of First Year Entering Students':'graduation_cohort','Number of Graduates':'graduation_completers','6-yr Graduation Rate':'graduation_rate'}[m]
    prefix='men_' if m.startswith('Gender_M_') else 'women_' if m.startswith('Gender_W_') else 'pell_' if m.startswith('Pell_Pell_') else 'nonpell_' if m.startswith('Pell_Non-Pell_') else 'nonresident_' if m.startswith('URM_International_') else None
    if prefix is None:return None
    return prefix+'graduation_'+('cohort' if m.endswith('_Cohort') else 'completers')


class PrototypeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.sample=execute(Plan(['applicants','admits','admission_rate','matriculants','yield_rate'],['Yale'],[2024]))

    def test_chat_result_controls_and_table_preserve_result(self):
        from streamlit.testing.v1 import AppTest
        from prototype.responses import finalize_answer
        result=execute(Plan(['graduation_rate'],['Yale'],list(range(2009,2019)),time_basis='entry_cohort'))
        answer=finalize_answer({'intent':'analyze','result':result,'plan':result['plan']},'Show Yale six-year graduation rate for entering cohorts 2009 through 2018')
        answer['model_usage']={'model':'test'}  # diagnostics must not render in chat
        app=AppTest.from_file(str(BASE/'prototype/app.py'),default_timeout=60).run()
        app.session_state['use_ai']=False  # Explicit offline path; AI routing has separate tests.
        app.session_state['history']=[{'question':'Graduation cohorts','answer':answer}]
        app.session_state['main_page']='Ask & Analyze'
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.expander),0)
        self.assertEqual(len(app.get('popover')),2)
        self.assertFalse(any(b.label in ['Show table only','Edit institutions, measures or years'] for b in app.button))
        toggle=next(t for t in app.toggle if t.label=='Display table')
        toggle.set_value(True).run()
        self.assertFalse(app.exception)
        self.assertTrue(len(app.dataframe)>0)
        next(c for c in app.checkbox if c.label=='Show underlying counts and data status').check().run()
        self.assertFalse(app.exception)
        numerical_tables=[table for table in app.dataframe if 'IPEDS release' not in table.value.columns]
        self.assertEqual(len(numerical_tables),2)
        self.assertTrue(any('status' in table.value.columns for table in numerical_tables))
        self.assertTrue(any('IPEDS release' in table.value.columns for table in app.dataframe))
        self.assertEqual(app.session_state['history'][0]['answer']['result'],result)

    def test_yale_2024(self):
        rows={r['metric']:r for r in self.sample['rows']}
        self.assertEqual(rows['applicants']['value'],57517)
        self.assertEqual(rows['admits']['value'],2227)
        self.assertEqual(rows['matriculants']['value'],1554)
        self.assertAlmostEqual(rows['admission_rate']['value'],100*2227/57517)
        self.assertAlmostEqual(rows['yield_rate']['value'],100*1554/2227)
        self.assertTrue(all(r['provisional'] for r in self.sample['rows']))
        self.assertTrue(self.sample['sources'][0]['codebook_evidence'])

    def test_entire_supported_benchmark(self):
        cases=json.loads((BASE/'evaluations/results/yale_cases.json').read_text(encoding='utf-8'))
        observed={}
        for metric in measures():
            result=execute(Plan([metric],['Yale'],list(range(2015,2025))))
            for row in result['rows']:observed[(metric,row['observation_year'])]=row
        counts={'pass':0,'mismatch':0}
        for case in cases:
            if case['status'] not in {'pass','mismatch'}:continue
            metric=fixture_metric(case)
            self.assertIsNotNone(metric,case['id'])
            row=observed[(metric,case['observation_year'])]
            self.assertEqual(row['status'],'ok',case['id'])
            value=row['value']/100 if row['unit']=='percent' else row['value']
            status='pass' if abs(value-float(case['expected']))<=float(case['tolerance'])+1e-12 else 'mismatch'
            self.assertEqual(status,case['status'],case['id'])
            self.assertAlmostEqual(value,float(case['actual']),places=10,msg=case['id'])
            counts[status]+=1
        self.assertEqual(counts,{'pass':229,'mismatch':13})
        type(self).benchmark_counts=counts

    def test_cohort_class_time_and_outside_coverage(self):
        wrong_model_plan={'metrics':['applicants'],'institutions':['Yale'],'years':[2024],'time_basis':'class','output':'table'}
        corrected=canonical_time('Yale class of 2028 entering fall 2024 applicants',wrong_model_plan)
        self.assertEqual(corrected['time_basis'],'observation');self.assertEqual(corrected['years'],[2024])
        rows=execute(Plan(['graduation_cohort','graduation_completers','graduation_rate'],['Yale'],[2009],'entry_cohort'))['rows']
        self.assertEqual([r['observation_year'] for r in rows],[2015]*3)
        self.assertEqual([r['value'] for r in rows[:2]],[1305,1263])
        self.assertEqual(execute(Plan(['applicants'],['Yale'],[2028],'class'))['rows'][0]['value'],57517)
        for plan in [Plan(['applicants'],['Yale'],[2025]),Plan(['graduation_rate'],['Yale'],[2019],'entry_cohort'),Plan(['applicants'],['Yale'],[2029],'class')]:
            with self.assertRaises(ValueError):execute(plan)

    def test_institutions_and_ivy(self):
        result=execute(Plan(['graduation_rate'],['Ivy League'],[2015]))
        self.assertEqual(len(result['rows']),8)
        self.assertTrue(all(r['status']=='ok' for r in result['rows']))
        self.assertEqual(institutions(['130794'],2024),[130794])
        with self.assertRaises(ValueError):institutions(['University'],2024)
        with self.assertRaises(ValueError):institutions(["Yale' OR 1=1 --"],2024)

    def test_uncollected_not_zero_and_retained_discrepancies(self):
        for metric,year in [('satvr50',2021),('pell_graduation_rate',2015)]:
            row=execute(Plan([metric],['Yale'],[year]))['rows'][0]
            self.assertEqual(row['status'],'not_collected');self.assertIsNone(row['value'])
        self.assertTrue(execute(Plan(['matriculants'],['Yale'],[2016]))['warnings'])
        for metric in ['urm_graduation_rate','satvr90']:
            with self.assertRaises(ValueError):execute(Plan([metric],['Yale'],[2024]))

    def test_degrees_aid_and_location_override(self):
        self.assertEqual(execute(Plan(['bachelors_degrees'],['Yale'],[2015]))['rows'][0]['value'],1327)
        aid=execute(Plan(['undergraduate_grant_percent'],['Yale'],[2015,2023]))
        row=aid['rows'][0];self.assertEqual((row['numerator'],row['denominator']),(2869,5477))
        self.assertAlmostEqual(row['value'],100*2869/5477)
        self.assertTrue(aid['sources'][1]['location_overrides'])

    def test_discovery_and_parser_followups(self):
        self.assertTrue(search('faculty',years=[2024]))
        self.assertTrue(search('SATVR25',years=[2024]))
        first=ask('Yale applicants and admits and admission rate in fall 2024')
        self.assertIn('admission_rate',first['plan']['metrics'])
        follow=ask('How many matriculated in that year and what is the yield rate?',first['plan'])
        self.assertEqual(follow['plan']['years'],[2024]);self.assertIn('yield_rate',follow['plan']['metrics'])
        self.assertEqual(local_plan('class of 2028 entering fall 2024 applicants for Yale')['plan']['years'],[2024])
        self.assertEqual(local_plan('What data do we have on graduate enrollment?')['intent'],'discover')
        self.assertEqual(local_plan('national graduation rate')['intent'],'catalog')
        for q in ['What percentage of Yale students were on financial aid?','Yale URM graduation rate','Yale SAT 90th percentile','Yale four-year graduation rate']:
            self.assertEqual(local_plan(q)['intent'],'clarify',q)
        with self.assertRaises(ValueError):ask('Yale applicants class of 2029')
        self.assertEqual(ask('Plot Yale six-year graduation rate from 2015 through 2024')['result']['coverage']['available_cells'],10)

    def test_exports_are_complete_and_attributable(self):
        from openpyxl import load_workbook
        wb=load_workbook(BytesIO(excel_bytes(self.sample)),read_only=True)
        self.assertEqual(wb['Data'].max_row,6)
        self.assertEqual(set(wb.sheetnames),{'Data','Methodology','Sources','Codebook evidence'})
        self.assertEqual(len(csv_bytes(self.sample).decode('utf-8-sig').splitlines()),6)
        self.assertEqual(json.loads(audit_bytes(self.sample))['result_id'],self.sample['result_id'])

    def test_read_only_and_plan_limits(self):
        with readonly(DATA/'metadata_catalog.sqlite') as db:
            with self.assertRaises(Exception):db.execute('CREATE TABLE prohibited(x)')
        for plan in [Plan(['applicants'],[],[2024]),Plan(['applicants'],['Yale'],[2024,2024]),Plan(['applicants'],['Yale'],[2024],'entry_cohort')]:
            with self.assertRaises(ValueError):plan.validate()
        for question in ['Yale graduate enrollment by race','Yale early decision admission rate','Yale women Pell graduation rate']:
            answer=ask(question)
            self.assertEqual(answer['intent'],'catalog_discovery')
            self.assertNotIn('result',answer)  # source planning is not numerical approval

    def test_streamlit_guided_chart_and_export(self):
        from streamlit.testing.v1 import AppTest
        app=AppTest.from_file(str(BASE/'prototype/app.py'),default_timeout=60).run()
        app.session_state['use_ai']=False  # Explicit offline path; AI routing has separate tests.
        app.session_state['development_mode']=True
        app.button(key='home_ask').click().run();app.button(key='open_builder').click().run()
        app.text_input(key='institution_query').set_value('Smith').run()
        self.assertTrue(len(app.multiselect(key='builder_inst_widget').options)>1)
        app.multiselect(key='builder_inst_widget').set_value([167835]).run()
        app.text_input(key='institution_query').set_value('Harvard').run()
        self.assertEqual(app.session_state['extract_ids'],[167835])
        app.multiselect(key='builder_inst_widget').set_value([167835,166027]).run()
        app.button(key='step_1').click().run()
        app.text_input(key='measure_query').set_value('APPLCN').run()
        app.button(key='measure_add_applicants').click().run()
        app.text_input(key='measure_query').set_value('admission rate').run()
        app.button(key='measure_add_admission_rate').click().run()
        self.assertEqual(set(app.session_state['extract_metrics']),{'applicants','admission_rate'})
        app.button(key='step_2').click().run()
        app.selectbox(key='extract_one_year_observation').set_value(2019).run()
        self.assertEqual(app.session_state['extract_years'],[2019])
        app.radio(key='builder_year_mode_widget').set_value('Selected years').run()
        app.multiselect(key='extract_selected_years_observation').set_value([2015,2018,2024]).run()
        self.assertEqual(app.session_state['extract_years'],[2015,2018,2024])
        app.radio(key='builder_year_mode_widget').set_value('Year range').run()
        app.selectbox(key='extract_from_observation').set_value(2023).run()
        app.selectbox(key='extract_through_observation').set_value(2024).run()
        app.button(key='step_3').click().run();app.button(key='preview_extract').click().run()
        self.assertFalse(app.exception)
        result=app.session_state['extract_result'];self.assertEqual(result['coverage']['available_cells'],8)
        rid=result['result_id'];app.button(key='builder_result_visualize').click().run()
        self.assertFalse(app.exception);self.assertEqual(app.session_state['extract_result']['result_id'],rid)
        self.assertTrue(app.session_state['chart_overrides'][rid])
        app.button(key='builder_back').click().run()
        self.assertEqual(len(app.session_state['history']),1)
        self.assertTrue(app.session_state['ideas_open'])

    def test_codebook_search_identifiers_and_full_definitions(self):
        records=search('APPLCN',[2024],40,kinds=['table_definition','variable_definition'])
        self.assertEqual((records[0]['table_name'],records[0]['variable_name']),('ADM2024','APPLCN'))
        self.assertTrue(all(r['kind'] in {'table_definition','variable_definition'} for r in records))
        detail=definition_detail(records[0]['record_id'])
        self.assertIn('first-time',detail['definition'].lower())
        self.assertNotIn('surveyorder:',detail['definition'])
        self.assertTrue(detail['sheet_name']);self.assertTrue(detail['source_row'])
        self.assertEqual(search('ADM2024',[2024])[0]['kind'],'table_definition')
        g=search('GRTYPE',[2015],40,kinds=['variable_definition'])[0]
        codes=definition_detail(g['record_id'])['codes']
        self.assertTrue(any(str(c['code_value'])=='8' for c in codes))
        self.assertEqual(search('zzzznonexistentidentifier',[2024]),[])

    def test_streamlit_codebook_search_submit_repeat_empty_and_categories(self):
        from streamlit.testing.v1 import AppTest
        app=AppTest.from_file(str(BASE/'prototype/app.py'),default_timeout=60).run()
        app.session_state['use_ai']=False  # Explicit offline path; AI routing has separate tests.
        app.session_state['development_mode']=True
        app.button(key='home_ask').click().run();app.button(key='open_definitions').click().run()
        app.text_input(key='codebook_search_query').set_value('APPLCN')
        app.button(key='FormSubmitter:codebook_search_form-Search codebooks').click().run()
        self.assertFalse(app.exception)
        records=app.session_state['codebook_search_result']['records']
        self.assertEqual(records[0]['variable_name'],'APPLCN')
        self.assertTrue(app.selectbox(key='codebook_browse_selected_record'))
        # A second submit replaces results; it must not collide with form/widget state.
        app.text_input(key='codebook_search_query').set_value('GRTYPE')
        app.multiselect(key='codebook_search_years').set_value([2015])
        app.button(key='FormSubmitter:codebook_search_form-Search codebooks').click().run()
        self.assertFalse(app.exception)
        self.assertEqual(app.session_state['codebook_search_result']['years'],[2015])
        self.assertTrue(any('Bachelor' in str(t.value) for t in app.dataframe))
        # No result and blank searches should clear old results, not leave a stale answer.
        app.text_input(key='codebook_search_query').set_value('zzzznonexistentidentifier')
        app.button(key='FormSubmitter:codebook_search_form-Search codebooks').click().run()
        self.assertFalse(app.exception)
        self.assertEqual(app.session_state['codebook_search_result']['records'],[])
        self.assertTrue(app.info)
        app.text_input(key='codebook_search_query').set_value(' ')
        app.button(key='FormSubmitter:codebook_search_form-Search codebooks').click().run()
        self.assertFalse(app.exception)
        self.assertTrue(app.info)

    def test_learning_examples_and_session_navigation(self):
        from prototype.workflows import topics,execute_example,compatible_wide,institution_matches
        from streamlit.testing.v1 import AppTest
        for topic in topics():
            self.assertEqual(len(topic['examples']),3)
            for ex in topic['examples']:
                answer=execute_example(ex['id'])
                self.assertEqual(answer['intent'],ex['kind'])
                if ex['kind']=='analyze':self.assertGreater(answer['result']['coverage']['available_cells'],0)
                else:self.assertTrue(answer['evidence'])
        self.assertIn(167835,institution_matches('Smith'))
        self.assertEqual(len(compatible_wide(self.sample)),1)
        with self.assertRaises(ValueError):compatible_wide(execute(Plan(['applicants','graduate_enrollment'],['Yale'],[2024])))
        app=AppTest.from_file(str(BASE/'prototype/app.py'),default_timeout=60).run()
        app.session_state['use_ai']=False  # Explicit offline path; AI routing has separate tests.
        app.button(key='topic_admissions').click().run();app.button(key='home_admissions-point').click().run()
        self.assertFalse(app.exception);self.assertEqual(len(app.session_state['history']),1)
        app.run();self.assertEqual(len(app.session_state['history']),1)  # no repeated execution
        app.button(key='nav_explore').click().run()
        app.button(key='topic_faculty').click().run();app.button(key='home_faculty-counts').click().run()
        self.assertEqual(len(app.session_state['history']),2)
        self.assertEqual(app.session_state['history'][-1]['answer']['intent'],'analyze')
        app.button(key='nav_about').click().run()
        app.selectbox(key='settings_model').set_value('claude-haiku-5-5').run()
        app.button(key='about_back').click().run();app.button(key='nav_about').click().run()
        self.assertEqual(app.session_state['model_choice'],'claude-haiku-5-5')
        self.assertFalse(app.exception)

    def test_topic_panel_row_placement_close_and_tab_navigation(self):
        from prototype.workflows import topics
        from streamlit.testing.v1 import AppTest
        app=AppTest.from_file(str(BASE/'prototype/app.py'),default_timeout=60).run()
        app.session_state['use_ai']=False  # Explicit offline path; AI routing has separate tests.
        self.assertFalse(app.radio)
        self.assertNotIn('close_topic',[b.key for b in app.button])
        content=topics()
        for topic in (content[0],content[3]):
            app.button(key='topic_'+topic['id']).click().run()
            self.assertFalse(app.exception)
            headings=[h.value for h in app.subheader]
            detail=topic['title']+' · '+topic['subtitle']
            detail_index=headings.index(detail)
            last_in_row=content[2] if topic==content[0] else content[5]
            self.assertGreater(detail_index,headings.index(last_in_row['icon']+' '+last_in_row['title']))
            if topic==content[0]:
                self.assertLess(detail_index,headings.index(content[3]['icon']+' '+content[3]['title']))
            self.assertEqual(sum(' · ' in h for h in headings),1)
            ex=topic['examples'][0]
            self.assertEqual(app.button(key='home_'+ex['id']).label,ex['question'])
            self.assertIn(ex['purpose'].upper(),[c.value for c in app.caption])
            app.button(key='close_topic').click().run()
            self.assertFalse(app.session_state['topic_open'])
            self.assertNotIn(detail,[h.value for h in app.subheader])
        app.button(key='nav_ask').click().run()
        self.assertEqual(app.session_state['main_page'],'Ask & Analyze')
        app.button(key='nav_explore').click().run()
        self.assertEqual(app.session_state['main_page'],'Explore IPEDS')
        self.assertFalse(app.exception)

    def test_scrollable_conversation_preserves_followups_and_tools(self):
        from streamlit.testing.v1 import AppTest
        app=AppTest.from_file(str(BASE/'prototype/app.py'),default_timeout=60).run()
        app.session_state['use_ai']=False  # Explicit offline path; AI routing has separate tests.
        app.button(key='topic_admissions').click().run()
        app.button(key='home_admissions-point').click().run()
        revision=app.session_state['transcript_revision']
        app.session_state['use_ai']=False
        app.chat_input(key='followup_question').set_value('How many applicants did Yale have in fall 2023?').run()
        self.assertEqual(app.session_state['transcript_revision'],revision+1)
        revision=app.session_state['transcript_revision']
        self.assertFalse(app.exception)
        self.assertEqual(len(app.session_state['history']),2)
        self.assertEqual(len(app.chat_message),4)
        self.assertIn('2023',app.session_state['history'][-1]['question'])
        self.assertTrue(app.session_state['ideas_open'])
        self.assertTrue(app.session_state['tools_open'])
        app.selectbox(key='ideas_topic_widget').set_value('enrollment').run()
        self.assertEqual(app.session_state['transcript_revision'],revision)
        app.button(key='collapse_ideas').click().run()
        app.button(key='collapse_tools').click().run()
        self.assertFalse(app.session_state['ideas_open'])
        self.assertFalse(app.session_state['tools_open'])
        self.assertEqual(len(app.chat_message),4)
        app.button(key='show_ideas').click().run()
        app.button(key='show_tools').click().run()
        self.assertEqual(app.selectbox(key='ideas_topic_widget').value,'enrollment')
        self.assertEqual(len(app.chat_message),4)
        app.session_state['development_mode']=True;app.run()
        app.button(key='open_builder').click().run()
        app.button(key='builder_back').click().run()
        self.assertEqual(len(app.chat_message),4)
        self.assertTrue(app.chat_input(key='followup_question'))
        app.button(key='new_conversation').click().run()
        self.assertEqual(len(app.chat_message),0)
        self.assertTrue(app.chat_input(key='initial_question'))
        app.session_state['use_ai']=False
        app.chat_input(key='initial_question').set_value('How many applicants did Yale have in fall 2024?').run()
        self.assertEqual(len(app.chat_message),2)
        self.assertTrue(app.chat_input(key='followup_question'))
        self.assertFalse(app.exception)

    def test_response_rules_summary_missing_context_and_ambiguity(self):
        from prototype.responses import summarize,presentation_for
        answer=ask('What is Yale admission rate in fall 2024?')
        self.assertIn('3.87%',answer['message'])
        self.assertNotIn('Using observation',answer['message'])
        self.assertEqual(answer['presentation']['chart_type'],'none')
        previous=ask('Plot Yale graduation rate from 2015 through 2024')['plan']
        fresh=ask('How many applicants were there in 2020?',previous)
        self.assertEqual(fresh['intent'],'clarify')  # no silent Yale inheritance
        fresh=ask('Harvard applicants in fall 2020',previous)
        self.assertEqual(fresh['plan']['years'],[2020])
        self.assertEqual(fresh['plan']['metrics'],['applicants'])
        fresh=ask('Show Harvard applicants',previous)
        self.assertEqual(fresh['plan']['years'],[2024])
        follow=ask('What about Harvard?',previous)
        self.assertIn('Interpreted', 'Interpreted as: '+follow['interpreted_as'])
        self.assertEqual(follow['plan']['metrics'],['graduation_rate'])
        cohort=ask('Yale graduation rate for students who started in 2009')
        self.assertEqual(cohort['result']['observation_years'],[2015])
        self.assertIn('entering cohort',cohort['message'])
        missing=execute(Plan(['satvr50'],['Yale'],[2021]))
        self.assertIn('not collected',summarize(missing))
        self.assertNotIn('suppressed',summarize(missing))
        self.assertNotIn('0%',summarize(missing))
        self.assertEqual(ask("Smith's graduation rate in 2024")['intent'],'clarify')
        self.assertEqual(presentation_for(execute(Plan(['admission_rate'],['Yale','Harvard'],[2024])))['chart_type'],'bar')
        self.assertEqual(presentation_for(execute(Plan(['admission_rate'],['Yale'],[2023,2024])))['chart_type'],'line')

    def test_display_edits_reuse_exact_result_and_preserve_unrequested_fields(self):
        from prototype.responses import finalize_answer,adjust_presentation
        answer=finalize_answer({'result':self.sample})
        edited=adjust_presentation('Set chart title to "Admissions comparison"',answer)
        self.assertIs(edited['result'],self.sample)
        self.assertEqual(edited['presentation']['title'],'Admissions comparison')
        edited=adjust_presentation('Make it a bar chart',edited)
        edited=adjust_presentation('Show me the table',edited)
        self.assertTrue(edited['presentation']['show_table'])
        self.assertEqual(edited['presentation']['title'],'Admissions comparison')
        self.assertEqual(edited['presentation']['chart_type'],'bar')
        self.assertEqual(edited['result']['result_id'],self.sample['result_id'])
        edited=adjust_presentation('Make the bars blue',edited)
        self.assertEqual(edited['presentation']['color'],'blue')
        self.assertEqual(adjust_presentation('Set y-axis to 0-200',finalize_answer({'result':execute(Plan(['admission_rate'],['Yale'],[2024]))}))['intent'],'clarify')
        self.assertEqual(adjust_presentation('Change the chart font',edited)['intent'],'clarify')
        self.assertEqual(adjust_presentation('Show me the table',None)['intent'],'clarify')
        self.assertIsNone(adjust_presentation('Show Harvard applicants in 2023',edited))

    def test_streamlit_chart_first_and_show_table_without_new_query(self):
        from unittest.mock import patch
        from streamlit.testing.v1 import AppTest
        app=AppTest.from_file(str(BASE/'prototype/app.py'),default_timeout=60).run()
        app.session_state['use_ai']=False  # Explicit offline path; AI routing has separate tests.
        app.button(key='topic_admissions').click().run()
        app.button(key='home_admissions-compare').click().run()
        self.assertFalse(app.exception)
        original=app.session_state['history'][-1]['answer']
        self.assertEqual(original['presentation']['chart_type'],'bar')
        self.assertFalse(original['presentation']['show_table'])
        # Collapsed detail tables remain available; main table is hidden for charts.
        self.assertTrue(app.get('vega_lite_chart'))
        with patch('prototype.ui.ask',side_effect=AssertionError('Display-only edits must not query')):
            app.chat_input(key='followup_question').set_value('Show me the table').run()
        self.assertFalse(app.exception)
        updated=app.session_state['history'][-1]['answer']
        self.assertEqual(updated['mode'],'presentation adjustment')
        self.assertEqual(original['result']['result_id'],updated['result']['result_id'])
        self.assertTrue(updated['presentation']['show_table'])

    def test_model_adapter_auto_strict_and_bounded_retry(self):
        from unittest.mock import patch,MagicMock
        from types import SimpleNamespace
        from prototype.agent import model_plan,MODEL_LABELS
        payload={'intent':'analyze','message':'Plan','query':'','metrics':['applicants'],
                 'institutions':['Yale'],'years':[2024],'time_basis':'observation','output':'table'}
        usage=SimpleNamespace(input_tokens=100,output_tokens=20,cache_creation_input_tokens=0,cache_read_input_tokens=0)
        text=SimpleNamespace(type='text',model_dump=lambda:{'type':'text','text':'Use a plan'})
        tool=SimpleNamespace(type='tool_use',name='choose_plan',input=payload)
        for model in MODEL_LABELS:
            with patch('anthropic.Anthropic') as cls:
                client=MagicMock();cls.return_value=client
                client.messages.create.side_effect=[SimpleNamespace(content=[text],usage=usage,stop_reason='end_turn'),SimpleNamespace(content=[tool],usage=usage,stop_reason='tool_use')]
                action=model_plan('Yale applicants fall 2024',{'metrics':['graduation_rate'],'institutions':['Harvard'],'years':[2015]},model=model)
                self.assertEqual(action['plan']['years'],[2024]);self.assertEqual(action['model_usage']['calls'],2)
                args=client.messages.create.call_args.kwargs
                self.assertIsNone(json.loads(args['messages'][0]['content'])['previous_plan'])
                self.assertEqual(args['tool_choice'],{'type':'auto'});self.assertTrue(args['tools'][0]['strict'])
                self.assertNotIn('maxItems',str(args['tools']))
        with patch('anthropic.Anthropic') as cls:
            cls.return_value.messages.create.return_value=SimpleNamespace(content=[text],usage=usage,stop_reason='end_turn')
            with self.assertRaises(ValueError):model_plan('Yale applicants fall 2024')
            self.assertEqual(cls.return_value.messages.create.call_count,2)
        self.assertEqual(canonical_time('grants for 2014–15 through 2023–24',{'metrics':['undergraduate_grant_percent'],'years':[2014,2023]})['years'],list(range(2015,2025)))


if __name__=='__main__':
    import datetime
    suite=unittest.defaultTestLoader.loadTestsFromTestCase(PrototypeTests)
    result=unittest.TextTestRunner(verbosity=2).run(suite)
    report={'timestamp':datetime.datetime.now(datetime.timezone.utc).isoformat(),
            'tests_run':result.testsRun,'failures':len(result.failures),'errors':len(result.errors),
            'benchmark_cells':getattr(PrototypeTests,'benchmark_counts',None),
            'scope':'Real-data deterministic retrieval, plan validation, local parser, full exports and Streamlit AppTest. Does not certify LLM intent accuracy or resolve source discrepancies.'}
    (BASE/'evaluations/prototype_verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    sys.exit(not result.wasSuccessful())


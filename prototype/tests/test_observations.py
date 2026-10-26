"""Typed categorical answers checked against source SQL, not planner output."""
from copy import deepcopy
from io import BytesIO
import json
import unittest
from unittest.mock import patch
from prototype.catalog import readonly, annual_path
from prototype.observations import answer, bind, decode, verify_result
from prototype.tests.test_query_engine import source


def projection(fields, institutions=('167835',), year=2024, filters=None, population=None):
    s,_=source('IC'+str(year),year,fields=fields,filters=filters)
    return {'version':1,'institutions':list(institutions),'population':population,
            'sources':[{k:s[k] for k in ('year','table','field_record_ids','filters')}],
            'result_kind':'institution_list' if population else 'lookup'}


POP={'fips':[25],'controls':[],'levels':[],'degree_granting':True,'title_iv':False,'administrative_offices':False}


class ObservationTests(unittest.TestCase):
    def test_three_school_services_and_source_codes(self):
        p=projection(['STUSRV2','STUSRV3','STUSRV4','STUSRV8'],('167835','166939','168218'))
        a=answer(p,'Compare student services in 2024-25')
        self.assertEqual(len(a['result']['rows']),12)
        for row in a['result']['rows']:
            self.assertIsNone(row['value'])
            self.assertEqual(row['response_label'],'Implied no' if row['measure'].startswith('On-campus') else 'Yes')
            self.assertTrue(row['category_record_ids'])
        self.assertNotIn('count',a['message'])
        self.assertEqual(a['presentation']['chart_type'],'none')
        from prototype.responses import adjust_presentation
        chart=adjust_presentation('Show a bar chart',a)
        self.assertEqual(chart['presentation']['chart_type'],'none')
        self.assertIn('no new calculation',chart['message'])

    def test_affiliation_summary_is_per_institution(self):
        a=answer(projection(['RELAFFIL'],('164924','167835')),'Compare religious affiliations in 2024-25')
        self.assertIn('Religious affiliation: Roman Catholic at Boston College',a['message'])
        self.assertIn('Religious affiliation: Not applicable at Smith College',a['message'])

    def test_explicit_no_and_award_levels(self):
        a=answer(projection(['OPENADMP'],('165112','167835')),'Open admission policies in 2024-25')
        self.assertEqual({r['unitid']:r['response_label'] for r in a['result']['rows']},{165112:'Yes',167835:'No'})
        a=answer(projection(['LEVEL5','LEVEL7'],('167835','166939','168218')),'Award levels in 2024-25')
        self.assertEqual(len(a['result']['rows']),6)
        self.assertEqual([r['response_label'] for r in a['result']['rows'] if r['unitid']==168218],['Yes','Implied no'])

    def test_institution_lists_match_independent_sql(self):
        for variables,filters,where,count in [(['DSTNUGC','DSTNUGP'],{'DSTNCED1':'1'},'i.DSTNCED1=1',83),(['SLO6','STUSRV8'],{'SLO6':'1','STUSRV8':'1'},'i.SLO6=1 and i.STUSRV8=1',14)]:
            p=projection(variables,(),filters=filters,population=POP)
            a=answer(p,'Which Massachusetts degree-granting institutions match these responses in 2024-25?')
            with readonly(annual_path(2024)) as db:
                expected={r[0] for r in db.execute("select h.UNITID from HD2024 h join IC2024 i on h.UNITID=i.UNITID where h.STABBR='MA' and h.DEGGRANT=1 and "+where)}
            actual={r['unitid'] for r in a['result']['rows']}
            self.assertEqual(actual,expected);self.assertEqual(len(actual),count)

    def test_cross_release_cells_and_missing_year_block(self):
        p=projection(['CREDITS3'])
        p['sources']+=projection(['CREDITS3'],year=2015)['sources']
        a=answer(p,'Did Smith accept AP credits in 2015-16 and 2024-25?')
        self.assertEqual([r['response_label'] for r in a['result']['rows']],['Yes','Yes'])
        p['sources']=p['sources'][:1]
        with self.assertRaisesRegex(ValueError,'every requested release'):bind(p,'AP credits in 2015-16 and 2024-25')

    def test_companion_bound_exact_percentage_and_observation_date(self):
        p=projection(['DISABPCT'],('167835','164739'))
        a=answer(p,'What disability-services percentage is in the 2024-25 release?')
        rows={r['unitid']:r for r in a['result']['rows']}
        self.assertEqual(rows[167835]['value'],38)
        self.assertEqual(rows[167835]['reporting_period'],'Fall 2023')
        self.assertEqual(rows[164739]['response_label'],'3 percent or less')
        self.assertEqual(rows[164739]['status'],'bounded')
        self.assertIsNone(rows[164739]['value'])
        self.assertEqual(rows[164739]['bound']['upper'],3)
        self.assertEqual(rows[164739]['reporting_period'],'Fall 2023')
        self.assertEqual(len(a['result']['sources'][0]['codebook_evidence']),2)
        with self.assertRaisesRegex(ValueError,'student observation year'):
            answer(p,'What percentage of undergraduates registered with disability services in Fall 2024?')
        # Persisted plans remain valid inputs, without internal server-only fields.
        again=answer(a['result']['observation_plan'],'Disability percentage in 2024-25')
        self.assertEqual(again['result']['rows'][0]['response_status'],a['result']['rows'][0]['response_status'])

    def test_mixed_values_exports_and_render(self):
        from prototype.exports import csv_bytes,excel_bytes,audit_bytes
        from openpyxl import load_workbook
        a=answer(projection(['CALSYS','DISABPCT']),'Calendar and disability percentage in 2024-25')
        r=a['result']
        self.assertIn('Semester',a['message']);self.assertIn('38.00%',a['message'])
        self.assertIn(b'raw_code',csv_bytes(r));self.assertIn(b'Semester',csv_bytes(r))
        wb=load_workbook(BytesIO(excel_bytes(r)))
        self.assertIn('Codebook evidence',wb.sheetnames)
        self.assertEqual(json.loads(audit_bytes(r))['rows'][0]['raw_code'],1)
        from streamlit.testing.v1 import AppTest
        def app(result):
            from prototype.components import render_result
            render_result(result,'typed_test',show_chart=True)
        app=AppTest.from_function(app,args=(r,)).run(timeout=60)
        self.assertFalse(app.exception)
        self.assertEqual(len(app.get('arrow_vega_lite_chart')),0)
        self.assertGreater(len(app.dataframe),0)

    def test_decode_negative_unknown_missing_and_conflicting_codes(self):
        codes=[{'record_id':i,'code_value':c,'value_label':l} for i,(c,l) in enumerate([('1','Yes'),('0','Implied no'),('-1','Not reported'),('-2','Not applicable')])]
        for raw,status in [(None,'missing'),('','missing'),(0,'implied_no'),(-1,'not_reported'),(-2,'not_applicable'),(999,'unmapped')]:
            self.assertEqual(decode(raw,codes)['response_status'],status)
        with self.assertRaises(ValueError):decode('1',codes+[{'record_id':55,'code_value':'1','value_label':'No'}])

    def test_types_roll_back_and_integrity(self):
        p=projection(['CALSYS']);a=answer(p,'Calendar in 2024-25')
        bad=deepcopy(a['result']);bad['rows'][0]['value']=1
        with self.assertRaises(ValueError):verify_result(bad)
        bad=deepcopy(a['result']);bad['rows']=[]
        with self.assertRaises(ValueError):verify_result(bad)
        with patch.dict('os.environ',{'IPEDS_TYPED_OBSERVATIONS':'0'}):
            with self.assertRaisesRegex(ValueError,'disabled'):bind(p,'Calendar in 2024-25')
        with self.assertRaisesRegex(ValueError,'numerical query'):bind(p,'How many institutions have this calendar in 2024-25?')
        with self.assertRaisesRegex(ValueError,'quality'):bind(p,'Which has better counseling in 2024-25?')

    def test_category_definition_and_blocked_followup(self):
        from prototype.institution_confirmation import ask
        for q in ['In the 2024-25 Institutional Characteristics codebook, what does “Implied no” mean for on-campus childcare, and how is it different from “Not reported”?','Define the label Implied no for on-campus childcare in the 2024-25 Institutional Characteristics codebook.']:
            with patch('prototype.institution_confirmation.mentions',side_effect=AssertionError('No model identity call')):
                a=ask(q,use_model=True)
            self.assertEqual(a['intent'],'definition')
            self.assertIn('0 = Implied no',a['message'])
        p=projection(['DSTNUGC','DSTNUGP'],('168218',))
        history=[{'question':'Distance courses or programs?','answer':{'observation_plan':p,'intent':'clarify'}}]
        a=ask('What is the difference between those two categories?',conversation=history)
        self.assertEqual(a['intent'],'definition')
        self.assertIn('all the required coursework',a['message'])

    def test_lossless_legacy_projection_and_arithmetic_stays_numeric(self):
        from prototype.observations import from_legacy
        s,_=source('IC2024',fields=['CALSYS'])
        from prototype.tests.test_query_engine import plan
        p=plan([s],[('fs_'+str(s['field_record_ids'][0]),'Calendar')],institutions=('167835',))
        adapted=from_legacy(p)
        self.assertEqual(answer(adapted,'Calendar in 2024-25')['result']['rows'][0]['response_label'],'Semester')
        p['expressions']=[{'id':'total','operator':'sum','inputs':[p['outputs'][0]['reference']],'group_by_references':[]}]
        self.assertIsNone(from_legacy(p))
        from prototype.query_engine import inspect
        p['outputs'][0]['reference']='xtotal'
        with self.assertRaises(ValueError):inspect(p,'Sum calendar codes in 2024-25')

    def test_requested_identity_and_source_concepts_cannot_drop(self):
        from prototype.intent_checks import CURRENT
        token=CURRENT.set({'unitids':[167835,166939]})
        try:
            with self.assertRaisesRegex(ValueError,'requested institutions'):bind(projection(['CALSYS']),'Calendar in 2024-25')
        finally:CURRENT.reset(token)
        with self.assertRaisesRegex(ValueError,'omits requested source concepts'):bind(projection(['STUSRV2']),'Compare counseling and on-campus childcare in 2024-25')

    def test_numerical_selector_dimension_does_not_change_executor(self):
        from prototype.observations import from_legacy
        s,ids=source('EF2024A',fields=['EFTOTLT','EFALEVEL'])
        p={'version':1,'institutions':['Yale'],'source_years':[2024],
           'selections':[{'record_id':ids['EFTOTLT'],'role':'measure'},{'record_id':ids['EFALEVEL'],'role':'dimension'}],
           'filters':[],'calculations':[],'joins':[],'clarification_questions':[]}
        self.assertIsNone(from_legacy(p))

    def test_mixed_calendar_and_direct_enrollment_with_source_selector(self):
        p=projection(['CALSYS'])
        s,_=source('EF2024A',fields=['EFTOTLT'],filters={'EFALEVEL':'2'})
        p['sources'].append({k:s[k] for k in ('year','table','field_record_ids','filters')})
        a=answer(p,'What calendar system does Smith College use and how many undergraduates does it have in the 2024-25 source release?')
        with readonly(annual_path(2024)) as db:
            expected=db.execute('select EFTOTLT from EF2024A where UNITID=167835 and EFALEVEL=2').fetchone()[0]
        rows=a['result']['rows']
        self.assertEqual(rows[0]['response_label'],'Semester')
        self.assertEqual(rows[1]['value'],float(expected))
        self.assertEqual(rows[1]['unit'],'students')
        self.assertTrue(a['result']['sources'][1]['code_filters'])

    def test_tool_batch_limit_recovers_with_usage(self):
        from prototype.tests.test_catalog_planner import block,response
        from prototype.catalog_planner import model_catalog_plan
        calls=[block('search_catalog',{'query':'calendar','year':2024,'limit':2},str(i)) for i in range(5)]
        with patch('anthropic.Anthropic') as provider:
            provider.return_value.messages.create.side_effect=[response(*calls),response(block('submit_observation_plan',projection(['CALSYS']),'final'))]
            a=model_catalog_plan('What calendar system did Smith College use in 2024-25?')
        self.assertEqual(a['result']['rows'][0]['response_label'],'Semester')
        self.assertEqual(a['model_usage']['calls'],2)
        second=provider.return_value.messages.create.call_args.kwargs['messages']
        self.assertTrue(any('Batch limit' in str(m) for m in second))


if __name__=='__main__':unittest.main()

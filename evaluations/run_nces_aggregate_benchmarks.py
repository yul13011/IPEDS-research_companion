"""Read-only NCES aggregate baseline: application coverage and independent source diagnostics.

Run from agent-redesign with its bundled Python. Never selects a mapping based on
agreement with expected values. Reference SQL is NOT the production query engine.
"""
from pathlib import Path
from collections import Counter
from decimal import Decimal
from unittest.mock import patch
import csv, datetime, hashlib, json, math, re, statistics, sys, time

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / '.packages'), str(ROOT)]
import openpyxl
from prototype.agent import ask
from prototype.catalog import annual_path, readonly
from analysis_catalog.catalog import connection

OUT = ROOT / 'evaluations' / 'nces_aggregates'
INPUT = ROOT / 'Test cases' / 'NCES'
PENDING = {
    8: 'Selected age category is not identified in the exported workbook.',
    200: 'Selected Outcome Measures cohort/population and outcome are not identified in the export.',
    110: 'Exact national core-expense/FTE aggregation and cross-component reporting-unit alignment need verification.',
    43: 'Exact NCES choice of estimated/reported FTE and professional-practice treatment needs verification.',
}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book.active
    rows = list(sheet.iter_rows(values_only=True))
    title = str(rows[2][0]).strip()
    values, notes = [], []
    for row, cells in enumerate(rows, 1):
        first = str(cells[0] or '').strip()
        second = cells[1] if len(cells) > 1 else None
        if re.fullmatch(r'\d{4}(?:-\d{2,4})?', first) and second is not None:
            values.append({'source_row': row, 'period': first, 'expected': float(second)})
        elif row > 5 and first and second is None:
            notes.append(first)
    result = {'q': int(re.search(r'-q(\d+)', path.name)[1]), 'file': str(path),
              'sha256': sha(path), 'sheet': sheet.title, 'title': title,
              'headers': list(rows[4]), 'notes': notes, 'observations': values,
              'sheets': book.sheetnames}
    book.close()
    return result


def source_year(q, period):
    start = int(period[:4])
    if q in {19, 20}: return start + 6
    if q in {21, 22}: return start + 3
    if q == 200: return start + 8
    if q in {2, 24, 25, 34, 35, 36, 37, 38, 43, 210, 220, 221}: return start + 1
    return start


def spec(q, y):
    """Candidate methods grounded in labels/codebooks; not fitted to workbook values."""
    ef = f'EF{y}'
    sfa = f'SFA{str(y-1)[2:]}{str(y)[2:]}' + ('_P1' if y<2024 else '')
    finance = f'F{str(y-1)[2:]}{str(y)[2:]}'
    # table, required columns, row selectors, operation, numerator, denominator,
    # post-reduction scale, extra directory filter
    if q == 1: return (f'HD{y}', [], {}, 'count', [], [], 1, '')
    if q in {3, 4, 5}: return (ef, ['EFTOTAL'], {'EFLEVEL': str({3:10,4:35,5:31}[q])}, 'sum', ['EFTOTAL'], [], 1, '')
    if q == 2: return (f'EFFY{y}', ['EFYTOTLT'], {'EFFYLEV':'1'}, 'sum', ['EFYTOTLT'], [], 1, '')
    if q in {150, 155}: return (ef+'C', ['EFRES01' if q==150 else 'EFRES02'], {'EFCSTATE':'99'}, 'sum', ['EFRES01' if q==150 else 'EFRES02'], [], 1, '')
    if q == 42: return (ef+'A_DIST', ['EFDEEXC','EFDESOM','EFDETOT'], {'EFDELEV':'1'}, 'ratio', ['EFDEEXC','EFDESOM'], ['EFDETOT'], 1, '')
    if q in {32, 33}:
        n,d = ('RET_NMF','RRFTCTA') if q==32 else ('RET_NMP','RRPTCTA')
        return (ef+'D',[n,d],{},'ratio',[n],[d],1,'')
    if q in {19,20,21}:
        return (f'GR{y}',['GRTOTLT'],{},'gr',[],[],1,'')
    if q == 22: return (f'GR{y}_L2',['LINE_11','LINE_50'],{},'ratio',['LINE_11'],['LINE_50'],1,'')
    if q == 24: return (f'C{y}_A',['CTOTALT'],{'CIPCODE':'99','MAJORNUM':'1','AWLEVEL':'15'},'sum',['CTOTALT'],[],1,'')
    if q == 210: return (f'C{y}_B',['CSTOTLT'],{},'sum',['CSTOTLT'],[],1,'')
    if q in {25,34,35,36,37,38,220,221}:
        n,d = {25:('ANYAIDN',None),34:('ANYAIDN','SCUGFFN'),35:('UPGRNTN','SCUGRAD'),36:('UPGRNTT','UPGRNTN'),37:('UFLOANN','SCUGRAD'),38:('UFLOANT','UFLOANN'),220:('AGRNT_T','AGRNT_N'),221:('LOAN_T','LOAN_N')}[q]
        return (sfa,[n]+([d] if d else []),{},'ratio' if d else 'sum',[n],[d] if d else [],1,'')
    if q in {101,102,103}:
        n,d = {101:('APPLCN',None),102:('ADMSSN','APPLCN'),103:('ENRLT','ADMSSN')}[q]
        return (f'ADM{y}',[n]+([d] if d else []),{},'ratio' if d else 'sum',[n],[d] if d else [],1,'')
    if q == 30: return (f'S{y}_OC',['HRTOTLT'],{'STAFFCAT':'1100'},'sum',['HRTOTLT'],[],1,'AND h.DEGGRANT=1')
    if q == 51: return (f'S{y}_IS',['HRTOTLT'],{'SISCAT':'1'},'sum',['HRTOTLT'],[],1,'AND h.DEGGRANT=1')
    if q in {12,50}:
        n,d = ('SAEQ9OT','SATOTLT') if q==50 else ('SATOTLT',None)
        return (f'SAL{y}_IS',[n]+([d] if d else []),{'ARANK':'7'},'ratio' if d else 'sum',[n],[d] if d else [],1,'AND h.DEGGRANT=1')
    if q in {13,15,17,39,40,41}:
        suffix,col,control = {13:('F1A','F1B25',1),15:('F2','F2B01',2),17:('F3','F3B01',3),39:('F1A','F1C191',1),40:('F2','F2E131',2),41:('F3','F3E071',3)}[q]
        return (finance+'_'+suffix,[col],{},'sum',[col],[],.001,f'AND h.CONTROL={control}')
    if q in {205,206,207,208}:
        n,fee,control = {205:('TUITION2','FEE2',1),206:('TUITION3','FEE3',None),207:('TUITION6','FEE6',1),208:('TUITION7','FEE7',None)}[q]
        return (f'Cost1_{y}' if y==2024 else f'IC{y}_AY',[n,fee],{},'mean',[n,fee],[],1,'AND h.RPTMTH IN (1,3) AND '+('h.CONTROL=1' if control else 'h.CONTROL IN (2,3)'))
    raise ValueError('No reviewed candidate mapping')


def reference(case, obs, metadata):
    q,y = case['q'],source_year(case['q'],obs['period'])
    r = {**obs, 'q':q,'title':case['title'],'source_file':case['file'],
         'source_sha256':case['sha256'],'source_sheet':case['sheet'],'source_year':y,
         'status':None,'reason':'','actual':None,'difference':None,'elapsed_seconds':0}
    fraction = 'rate' in case['headers'][1].lower() or 'percent' in case['headers'][1].lower()
    averaged = 'average' in case['headers'][1].lower()
    r['tolerance'] = .0005 if fraction else .5 if averaged or q in {13,15,17,39,40,41,110} else 0
    r['tolerance_basis'] = '0.05 percentage points (benchmark precision)' if fraction else 'nearest displayed dollar/thousand-dollar unit' if r['tolerance'] else 'exact count'
    if q in PENDING:
        r.update(status='definition_pending',reason=PENDING[q]); return r
    if not 2015<=y<=2024:
        r.update(status='outside_local_release_coverage',reason='Required source release is outside local 2015–2024 annual databases.');return r
    start=time.perf_counter()
    try:
        table,cols,filters,op,nums,dens,scale,extra=spec(q,y)
        path=annual_path(y);r['database']=str(path)
        with readonly(path) as db:
            names={x[0].casefold():x[0] for x in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            if table.casefold() not in names:raise LookupError('Mapped source table unavailable: '+table)
            table=names[table.casefold()]
            physical={x[1].upper():x[1] for x in db.execute('PRAGMA table_info("'+table+'")')}
            # EFFYALEV replaced the older EFFYLEV when attendance detail was added.
            if q==2 and 'EFFYALEV' in physical: filters={'EFFYALEV':'1'}
            used=list(dict.fromkeys(cols+list(filters)+(['GRTYPE'] if op=='gr' else [])))
            if any(c.upper() not in physical for c in used):raise LookupError('Mapped source column unavailable: '+str([c for c in used if c.upper() not in physical]))
            defs={x['variable_name'].upper():dict(x) for x in metadata.execute('SELECT * FROM source.variable_docs WHERE release_id LIKE ? AND table_name=? COLLATE NOCASE',(f'IPEDS_{y}%',table))}
            if any(c.upper() not in defs for c in used):
                absent=[c for c in used if c.upper() not in defs]
                elsewhere=[dict(x) for c in absent for x in metadata.execute('SELECT record_id,table_name,variable_name FROM source.variable_docs WHERE release_id LIKE ? AND variable_name=? COLLATE NOCASE',(f'IPEDS_{y}%',c))]
                r['alternative_documentation_locations']=elsewhere
                raise LookupError('Required field definition absent for physical table '+table+': '+', '.join(absent)+'. Alternative documentation locations retained; no implicit ownership substitution.')
            # Read the exact directory universe, plus all source selectors, into the audit.
            selector_codes={c:[dict(x) for x in metadata.execute('SELECT record_id,code_value,value_label FROM source.code_values WHERE release_id LIKE ? AND table_name=? COLLATE NOCASE AND variable_name=? COLLATE NOCASE',(f'IPEDS_{y}%',table,c)) if x['code_value']==v] for c,v in filters.items()}
            if any(not v for v in selector_codes.values()):raise LookupError('Source selector lacks a documented code label')
            select=['s.UNITID AS UNITID']+[f's."{c}" AS "{c}"' for c in used]
            where=['h.FIPS BETWEEN 1 AND 56','h.PSET4FLG IN (1,3)']
            params=[]
            for c,v in filters.items():where.append(f's."{c}"=?');params.append(v)
            if op=='gr':where.append('s.GRTYPE IN (?,?)');params.extend({19:['3','2'],20:['12','8'],21:['30','29']}[q])
            sql='SELECT '+','.join(select)+f' FROM "{table}" s JOIN "HD{y}" h ON s.UNITID=h.UNITID WHERE '+' AND '.join(where)+' '+extra
            raw=[dict(x) for x in db.execute(sql,params)]
            keys=[(x['UNITID'],x.get('GRTYPE')) for x in raw]
            if len(keys)!=len(set(keys)):raise LookupError('Unexpected duplicate reporting-unit/selector rows')
            missing=Counter(); complete=[]
            from prototype.query_engine import number
            codes={c:{x['code_value'] for x in metadata.execute('SELECT code_value,value_label FROM source.code_values WHERE release_id LIKE ? AND table_name=? COLLATE NOCASE AND variable_name=? COLLATE NOCASE',(f'IPEDS_{y}%',table,c)) if re.search('not applicable|missing|suppressed|not reported|not available',x['value_label'],re.I)} for c in cols}
            for row in raw:
                item={c:number(row[c],defs[c],codes[c]) for c in cols}
                for c,v in item.items():
                    if v is None:missing[c]+=1
                if all(v is not None for v in item.values()):complete.append((row,item))
            if op=='count': actual=len({x['UNITID'] for x in raw});numerator=actual;denominator=None
            elif op=='gr':
                nc,dc={19:('3','2'),20:('12','8'),21:('30','29')}[q]
                paired={}
                for row,item in complete:paired.setdefault(row['UNITID'],{})[str(row['GRTYPE'])]=item['GRTOTLT']
                valid=[v for v in paired.values() if nc in v and dc in v]
                numerator=sum(Decimal(str(v[nc])) for v in valid);denominator=sum(Decimal(str(v[dc])) for v in valid)
                actual=float(numerator/denominator) if denominator else None
                r['unpaired_institutions']=len(paired)-len(valid)
            else:
                numerator=sum((sum(Decimal(str(item[c])) for c in nums) for _,item in complete),Decimal(0))
                denominator=(sum((sum(Decimal(str(item[c])) for c in dens) for _,item in complete),Decimal(0)) if op=='ratio' else len(complete) if op=='mean' else None)
                actual=float(numerator/Decimal(denominator)*Decimal(str(scale))) if denominator else float(numerator*Decimal(str(scale))) if op=='sum' else None
            r.update(actual=actual, sql=sql,parameters=params,operation=op,numerator=float(numerator),denominator=float(denominator) if denominator is not None else None,scale=scale,
                selected_source_rows=len(raw),selected_institutions=len({x['UNITID'] for x in raw}),complete_rows=len(complete),missing_cells_by_field=dict(missing),
                missing_policy='Complete pairs/rows only for reference diagnostics; missing cells are never zero. This policy is not approved for the production national executor.',
                field_evidence=[{k:defs[c][k] for k in ['record_id','variable_name','title','long_description']} for c in used],selector_evidence=selector_codes,
                universe='US 50 states + DC (FIPS 1–56), PSET4FLG 1 or 3; source-table reporting units; no fabricated national unique-person count.')
            if actual is None:r.update(status='no_usable_values',reason='No complete values or zero denominator.')
            else:
                diff=actual-obs['expected'];r['difference']=diff
                r['status']='match' if abs(diff)<=r['tolerance']+1e-10 else 'mismatch'
                r['reason']='Matches within displayed benchmark precision.' if r['status']=='match' else 'Candidate reference differs from NCES; cause unresolved (definition, universe, release revision, or calculation). Not automatically attributed to an app arithmetic error.'
    except LookupError as e:r.update(status='mapping_pending',reason=str(e))
    except Exception as e:r.update(status='execution_error',reason=f'{type(e).__name__}: {e}')
    r['elapsed_seconds']=time.perf_counter()-start
    return r


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    cases=[load(p) for p in sorted(INPUT.glob('*.xlsx'),key=lambda p:int(re.search(r'-q(\d+)',p.name)[1]))]
    app=[];results=[]
    with connection() as metadata:
        for case in cases:
            question='National aggregate for the United States: '+case['title']+'. Periods: '+', '.join(o['period'] for o in case['observations'])+'.'
            # Real analytical entry point, model enabled. Fail loudly if the known
            # scope guard unexpectedly falls through to any provider/planner call.
            start=time.perf_counter()
            with patch('prototype.agent.model_plan',side_effect=AssertionError('Unexpected model call')), patch('prototype.catalog_planner.model_catalog_plan',side_effect=AssertionError('Unexpected catalog model call')):
                answer=ask(question,use_model=True)
            app.append({'q':case['q'],'question':question,'intent':answer.get('intent'),'message':answer.get('message'),
                'status':'blocked_national_scope' if 'aggregation is deferred' in answer.get('message','') and not answer.get('result') else 'unexpected_response',
                'wall_seconds':time.perf_counter()-start,'latency':answer.get('latency'),'estimated_api_usd':0,'model_calls':0,
                'entry_point':'prototype.agent.ask, fresh context, use_model=True; production analytical route, without browser/institution-name UI stage'})
            results.extend(reference(case,obs,metadata) for obs in case['observations'])
            print(json.dumps({'q':case['q'],'application':app[-1]['status'],'reference':dict(Counter(r['status'] for r in results if r['q']==case['q']))}),flush=True)
    for name,obj in [('source_manifest.json',cases),('reference_results.json',results),('application_results.json',app)]:
        (OUT/name).write_text(json.dumps(obj,indent=2,ensure_ascii=False),encoding='utf-8')
    summary={'timestamp':datetime.datetime.now(datetime.timezone.utc).isoformat(),'files':len(cases),'observations':len(results),
        'application_status':dict(Counter(c['status'] for c in app)),'reference_status':dict(Counter(r['status'] for r in results)),
        'application_median_seconds':statistics.median(x['wall_seconds'] for x in app),'estimated_api_usd':0,
        'reference_median_seconds':statistics.median(r['elapsed_seconds'] for r in results if r['elapsed_seconds']),
        'scope':'Application guard coverage and separate read-only SQL reference diagnostics, not successful production national calculations or live AI planning latency.',
        'case_status':[{'q':c['q'],'title':c['title'],'counts':dict(Counter(r['status'] for r in results if r['q']==c['q']))} for c in cases]}
    (OUT/'summary.json').write_text(json.dumps(summary,indent=2,ensure_ascii=False),encoding='utf-8')
    fields=['q','title','source_file','source_sheet','source_row','period','source_year','expected','actual','difference','tolerance','status','reason','elapsed_seconds','selected_institutions','selected_source_rows','complete_rows','numerator','denominator','database','sql','parameters']
    with (OUT/'reference_results.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');writer.writeheader();writer.writerows(results)
    print(json.dumps(summary,indent=2),flush=True)


if __name__=='__main__':main()

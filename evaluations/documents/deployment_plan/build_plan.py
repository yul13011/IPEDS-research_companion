from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'IPEDS App Deployment Plan.docx'
d=Document();sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.75)
sec.header_distance=sec.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Header','Footer']:
    s=d.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
normal=d.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',23),('Heading 1',16),('Heading 2',12)]:
    s=d.styles[name];s.font.size=Pt(size);s.font.bold=True
    s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(6)
    s.paragraph_format.keep_with_next=True
d.styles['Title'].paragraph_format.space_before=Pt(0)
d.styles['Subtitle'].font.size=Pt(11)
for style in d.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
h=sec.header.paragraphs[0];h.text='IPEDS RESEARCH COMPANION';h.runs[0].font.size=Pt(8)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('Deployment plan  |  October 26 2026  |  Page ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)

def p(text='',style=None):return d.add_paragraph(text,style)
def h1(text):d.add_heading(text,1)
def h2(text):d.add_heading(text,2)
def bullet(text):p(text,'List Bullet')
def page(title):d.add_page_break();h1(title)
def link(label,url):
    par=p();rid=par.part.relate_to(url,RT.HYPERLINK,is_external=True)
    node=OxmlElement('w:hyperlink');node.set(qn('r:id'),rid)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'000000');props.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');props.append(underline);run.append(props)
    text=OxmlElement('w:t');text.text=label;run.append(text);node.append(run);par._p.append(node)
def table(headers,rows,widths,size=10.5):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,v in zip(t.rows[0].cells,headers):c.text=v
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for i,row in enumerate(t.rows):
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'333333' if i==0 else 'F5F5F5' if i%2==0 else 'FFFFFF');cp.append(sh)
            bd=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
            cp.append(bd);m=OxmlElement('w:tcMar')
            for edge,val in [('top',85),('bottom',85),('left',95),('right',95)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');m.append(e)
            cp.append(m)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.0
                for run in par.runs:
                    run.font.size=Pt(size);run.font.bold=i==0;run.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
    p().paragraph_format.space_after=Pt(0)
    return t

p('IPEDS App Deployment Plan','Title')
p('Redesigned app  |  Working plan  |  October 26 2026','Subtitle')
p('This document records how we can deploy the redesigned IPEDS Research Companion for other people to use. It preserves the implementation sequence, launch checks and unresolved concerns so we can resume without reconstructing the earlier discussion.')
p('Recommendation: use a new GitHub repository, Streamlit Community Cloud and a private Cloudflare R2 Standard bucket for a limited pilot. Keep the existing graduation-rate prototype and its URL unchanged. The design is feasible; cloud implementation and hosted capacity testing are still pending.')
h2('Current status')
p('The owner has GitHub and Streamlit accounts. The local code and data were reviewed, and current hosting documentation was checked. R2 configuration, uploads, a clean Linux build and remote deployment have not been completed. This plan does not establish production capacity or set a public launch date.')
h2('Proposed services and initial costs')
table(['Service','Purpose','Expected initial cost'],[
    ['GitHub','Code, configuration, small assets and data manifest','Existing account'],
    ['Streamlit Community Cloud','Runs the interface, SQLite queries and Claude requests','Free hosting [1]'],
    ['R2 Standard','Durable copies of the versioned database release','Potentially $0 within allowances [2]'],
    ['Anthropic API','Prompt interpretation and source-backed planning','Usage based; set a spending limit [3]']], [1.5,3.55,1.95])
h2('Data and operating model')
p('The candidate runtime release is 13 SQLite databases totaling 3,057,135,616 bytes: 3.057 GB or 2.847 GiB. It contains ten annual files plus metadata, analysis and unit catalogs. Keep these large files outside ordinary Git source control.')
p('The Streamlit server downloads the files from R2, verifies them and queries local copies. R2 stores files; it does not execute the SQLite queries. Each visitor receives answers, charts and requested downloads rather than the full database package.')
p('R2 currently includes 10 GB-month Standard storage, one million Class A and ten million Class B operations monthly, with free direct egress. The current release fits if other account usage leaves room. Storage above the allowance is $0.015 per GB-month; requests can add charges. An R2 subscription must still be enabled. Recheck prices before setup. [2]')

page('Deployment implementation')
h2('Phase 1 Prepare a separate release package')
p('Create a clean repository for the redesign. The current local workspace still points to the original graduation-rate Git remote, so do not push the entire workspace into that deployment.')
bullet('Include prototype, analysis_catalog, the reviewed measure registry, required JSON and JavaScript assets, root Streamlit configuration and declared dependencies. Use prototype/app.py as the entry point.')
bullet('Exclude local .packages, environments, credentials, raw Access files, conversion outputs, proposal logs and test artifacts. Preserve the existing relative directory layout.')
bullet('Create a manifest of database object keys, relative paths, sizes and SHA-256 checksums. Pair the code revision with one specific data release.')
p('Output: a reviewed code bundle and manifest. Static inspection suggests field_semantics.sqlite is a preparation artifact rather than a runtime dependency; confirm its exclusion in a clean-package smoke test. Nothing should be deleted as part of packaging.')
h2('Phase 2 Configure R2 and publish the data release')
bullet('Enable R2 Standard storage and create a private bucket. Upload the 13 reviewed runtime databases using versioned object paths.')
bullet('Use separate credentials for uploading and app access. Give the hosted app bucket-scoped Object Read only access through the S3-compatible API. [4]')
bullet('Record the bucket, release identifier and manifest location. Keep credentials in hosted secrets, never in GitHub or browser code.')
h2('Phase 3 Add reliable data initialization')
p('Add the downloader before the UI checks for its local metadata catalog. There is no R2 downloader in the current app. Add and pin the selected Python S3 client dependency.')
bullet('Stream downloads to disk, bound parallel downloads and guard initialization across simultaneous sessions. Use temporary files, checksum verification and atomic completion.')
bullet('Reuse a completed matching release across sessions and reruns. Show clear progress and recover from interrupted downloads or missing files without serving mismatched catalogs.')
bullet('Start with the complete data package. Current coverage and guide inventory inspect all annual sources; fetching just one year would require more changes.')
h2('Phase 4 Configure hosted execution and controls')
p('Use Python 3.12 as the initial clean-build target, matching the local major/minor version. Install declared packages without Windows binaries or OneDrive paths. Verify Linux file-name case, SQLite full-text support and chart/image exports.')
p('Set server-side Anthropic and R2 secrets. Use a dedicated Anthropic workspace/key with spend and rate limits. Add application request limits, bounded concurrent model work and an operator AI cutoff. Keep Automatic as the public model default: Haiku for interpretation/reviewed planning, Sonnet for catalog planning/writing. Restrict development tools and model overrides to server configuration or authorized operators. [3, 5]')

page('Validation and launch')
h2('Phase 5 Deploy a preview and measure it')
p('Connect the new repository to a separate Streamlit Community Cloud app, configure secrets and deploy a preview URL. Keep the original app and credentials separate. Start with one or two active users, then test five active sessions; these are test levels, not established capacity commitments.')
bullet('Measure first deployment, first data download, warm requests, restart and wake behavior. Verify that completed files are reused and missing files can be restored.')
bullet('Measure actual Linux/container memory and CPU, request latency, model calls, token cost and errors. Include long conversations, national analyses and simultaneous Excel downloads.')
bullet('Check general data/year questions, Ivy comparisons, unavailable Fall 2025 requests, definitions, categorical lookups, graduate percentages, mixed requests, chart edits and source citations.')
bullet('Verify session separation, model/API failures, rate-limit messages, refresh behavior and handling of interrupted downloads or invalid checksums.')
h2('Evidence already available')
p('The October 23 Windows probe observed approximately 202 MiB peak resident memory and 539 MiB sampled private committed memory. These measures are different and must not be added. Two retained UI test sessions plus simultaneous backend workloads completed, but this was not a two-browser hosted load test and included no Claude calls or R2 transfer.')
p('That probe predates the latest conversation, guide and task-tracking features. It supports trying a small pilot, while current hosted measurements must determine operating capacity. The data package is held on disk rather than loaded wholesale into RAM.')
p('Streamlit currently lists approximate resource limits of 0.078–2 CPU cores, 690 MB–2.7 GB RAM and up to 50 GB disk, with figures still labeled February 2024. They are variable limits rather than reserved resources. Apps hibernate after 12 hours without traffic. Local-file persistence is not guaranteed. [6, 7]')
h2('Phase 6 Release a limited public pilot')
p('Release only after the preview reproduces expected answers, verifies every database, respects usage controls and passes the agreed workload. Decide acceptable response times, concurrency and spending before judging these tests. Keep known unsupported calculations and benchmark limitations visible.')
bullet('Record the deployed code revision, data release, dependency versions, operator settings and launch checks.')
bullet('Explain that chats are temporary, which information is sent to Anthropic and whether any questions are retained. Do not add durable question logging without a retention decision.')
bullet('Review costs and errors during the pilot before expanding access. Claude usage is expected to be the main variable expense; there is no dependable fixed price per question yet.')
h2('Rollback and fallback hosting')
p('Retain the previous code revision and matching data manifest. Roll back both together; disable new AI requests if cost or reliability fails. If free hosting cannot sustain the measured workload or startup time, use a host with explicitly allocated CPU/RAM and persistent disk. SQLite and R2 can remain in the architecture.')

page('Areas of concern')
p('Resolve these items before the relevant launch stage. Their importance is based on the current implementation and deployment design; the list is not evidence that any failure has occurred on Streamlit.')
table(['Concern','Why it matters','Response and verification'],[
 ['Cold starts and storage','About 3 GB may need restoring. Host-local files are not durable.','Measure initial download and wake/restart behavior. Cache verified files, restore from R2 and recover partial transfers.'],
 ['Shared memory and CPU','Free-host capacity varies. National queries, large histories and concurrent exports can raise peaks.','Test current code with real sessions. Bound work and monitor usage; move hosts if measured headroom is inadequate.'],
 ['Public API spending','Visitors share a paid key. Session limits can be bypassed by recreating sessions.','Set provider workspace spend/rate limits, app request limits, bounded concurrency and an operator cutoff.'],
 ['Developer controls','The current About page lets every visitor enable tools or select a more expensive model.','Treat the switch as a UI convenience, not authorization. Enforce public settings on the server; test visitor access.'],
 ['Release consistency','A partial or mixed release can invalidate catalogs or produce incorrect bindings.','Use a versioned checksum manifest, atomic completion and code/data pairing. Fail clearly on mismatch.'],
 ['User information','Proposal JSON can contain question/context data; local files can accumulate or disappear.','Choose temporary versus durable retention. Bound writes, keep credentials out of logs and document any collection.'],
 ['Linux portability','The current development environment uses Windows and project-local installed packages.','Run a clean dependency install and source/export smoke tests without local binaries or absolute Windows paths.'],
 ['Analytical reliability','Passing small tests does not establish accuracy for arbitrary IPEDS questions.','Retain source, population and formula validation, explicit uncertainty and unsupported-query safeguards. Expand pilot cases.']], [1.35,2.35,3.3],10.5)
h2('What is not established yet')
p('No Linux cloud build, R2 transfer benchmark, deployed memory measurement or hosted concurrency test has been completed. Hosting feasibility should not be presented as a guarantee of uptime, user capacity or universal analytical accuracy.')

page('Decisions and where to resume')
h2('Decisions still needed')
bullet('Pilot audience: invited users first, or anonymous public access with appropriate usage controls.')
bullet('Spending: monthly Anthropic cap, request allowance, concurrency limit and who monitors costs and errors.')
bullet('Ownership: new GitHub repository, Streamlit app name/URL and R2 bucket under the intended accounts.')
bullet('Retention: whether proposals, feedback or questions should be kept, for how long and in which durable store.')
bullet('Acceptance targets: acceptable cold-start time, warm response time and concurrent workload before public launch.')
h2('Resume from this point')
p('Next implementation task: prepare the separate code bundle and data manifest, then build the R2 bootstrap and public-use controls. Account setup and upload follow the reviewed release package. No application or hosting configuration was changed during the feasibility investigation.')
p('Project folder: C:/Users/liyuj/OneDrive/2026/IPEDS Project/agent-redesign')
p('Supporting files below are relative to that folder. This Word plan records the October 26 position; older deployment notes are historical context.')
for text in ['DEPLOYMENT_FEASIBILITY_2026-10-26.md — detailed investigation and source links.',
 'evaluations/deployment_feasibility_inventory.json — exact candidate database paths and sizes.',
 'LOCAL_PERFORMANCE_REPORT.md — October 23 local test method, results and limits.',
 'evaluations/capacity_local_results.json — underlying local performance observations.',
 'DEPLOYMENT_PLAN.md and STREAMLIT_CAPACITY_REVIEW.md — earlier planning notes.']:
    bullet(text)
h2('Official references checked October 26 2026')
refs=[('1 Streamlit Community Cloud','https://docs.streamlit.io/deploy/streamlit-community-cloud'),
 ('2 Cloudflare R2 pricing','https://developers.cloudflare.com/r2/pricing/'),
 ('3 Anthropic workspace spending and rate controls','https://platform.claude.com/docs/en/manage-claude/workspaces'),
 ('4 R2 authentication and token permissions','https://developers.cloudflare.com/r2/api/tokens/'),
 ('5 Streamlit hosted secrets','https://docs.streamlit.io/deploy/concepts/secrets'),
 ('6 Streamlit resources and hibernation','https://docs.streamlit.io/deploy/streamlit-community-cloud/manage-your-app'),
 ('7 Streamlit local storage qualification','https://docs.streamlit.io/develop/concepts/connections/connecting-to-data'),
 ('8 Streamlit repository and dependency layout','https://docs.streamlit.io/deploy/streamlit-community-cloud/deploy-your-app/file-organization')]
for label,url in refs:
    link(label,url)
    d.paragraphs[-1].paragraph_format.space_after=Pt(3)
p('Recheck service prices, allowances and hosting limits when resuming. They can change after this review.').runs[0].italic=True
d.core_properties.title='IPEDS App Deployment Plan'
d.core_properties.subject='Deployment sequence concerns and decisions for the redesigned IPEDS app'
d.core_properties.author=''
d.save(OUT)
print(OUT)

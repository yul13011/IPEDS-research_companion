"""Import original annual metadata faithfully and reconcile physical copies.

No measure definitions are harmonized and no question-bank tests are run.
"""
from pathlib import Path
import argparse,csv,hashlib,json,re,sqlite3,zipfile,xml.etree.ElementTree as ET
import openpyxl

BASE=Path(__file__).resolve().parents[1];PROJECT=BASE.parent;ROOT=PROJECT/'IPEDS Data';DATA=BASE/'data';TARGET=DATA/'metadata_catalog.sqlite'
def sha(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
    return h.hexdigest()
def val(v):return '' if v is None else str(v)
def clean(v):return re.sub(r'\s+',' ',val(v).replace('_x000D_',' ')).strip()
def js(v):return json.dumps(v,ensure_ascii=False,default=str)
DDL='''
PRAGMA foreign_keys=ON;
CREATE TABLE releases(release_id TEXT PRIMARY KEY,collection_cycle TEXT NOT NULL,release_label TEXT NOT NULL,release_label_basis TEXT NOT NULL,source_directory TEXT NOT NULL);
CREATE TABLE source_files(source_id INTEGER PRIMARY KEY,release_id TEXT REFERENCES releases,role TEXT,path TEXT UNIQUE,bytes INTEGER,sha256 TEXT,unchanged_after_read INTEGER);
CREATE TABLE sheets(sheet_id INTEGER PRIMARY KEY,source_id INTEGER REFERENCES source_files,sheet_name TEXT,role TEXT,max_row INTEGER,max_column INTEGER);
CREATE TABLE sheet_rows(sheet_id INTEGER REFERENCES sheets,row_number INTEGER,values_json TEXT NOT NULL,PRIMARY KEY(sheet_id,row_number));
CREATE TABLE records(record_id INTEGER PRIMARY KEY,release_id TEXT REFERENCES releases,source_id INTEGER REFERENCES source_files,sheet_id INTEGER REFERENCES sheets,source_row INTEGER,kind TEXT,table_name TEXT,variable_name TEXT,title TEXT,body TEXT,raw_json TEXT);
CREATE TABLE table_docs(record_id INTEGER PRIMARY KEY REFERENCES records,release_id TEXT,table_name TEXT,table_number TEXT,survey TEXT,title TEXT,description TEXT,year_coverage TEXT,release_status TEXT,release_date TEXT);
CREATE TABLE variable_docs(record_id INTEGER PRIMARY KEY REFERENCES records,release_id TEXT,table_name TEXT,variable_name TEXT,variable_number TEXT,title TEXT,long_description TEXT,data_type TEXT,format TEXT,field_width TEXT,multi_record TEXT,imputation_variable TEXT,variable_source TEXT,file_title TEXT,section_title TEXT);
CREATE TABLE code_values(record_id INTEGER PRIMARY KEY REFERENCES records,release_id TEXT,table_name TEXT,variable_name TEXT,variable_number TEXT,code_value TEXT,value_label TEXT,frequency TEXT,percent TEXT,value_order TEXT);
CREATE TABLE release_notes(record_id INTEGER PRIMARY KEY REFERENCES records,release_id TEXT,paragraph_number INTEGER,text TEXT);
CREATE VIRTUAL TABLE search_index USING fts5(release_id UNINDEXED,kind UNINDEXED,table_name,variable_name,title,body,record_id UNINDEXED,tokenize='unicode61');
CREATE TABLE physical_tables(release_id TEXT REFERENCES releases,table_name TEXT COLLATE NOCASE,source_rows INTEGER,sqlite_rows INTEGER,sqlite_path TEXT,integrity_verified INTEGER,values_verified INTEGER,PRIMARY KEY(release_id,table_name));
CREATE TABLE physical_columns(release_id TEXT,table_name TEXT COLLATE NOCASE,column_name TEXT COLLATE NOCASE,ordinal INTEGER,access_type TEXT,access_width INTEGER,sqlite_type TEXT,null_count INTEGER,PRIMARY KEY(release_id,table_name,column_name),FOREIGN KEY(release_id,table_name) REFERENCES physical_tables);
CREATE TABLE issues(issue_id INTEGER PRIMARY KEY,release_id TEXT,issue_type TEXT,severity TEXT,source_record_id INTEGER REFERENCES records,table_name TEXT,variable_name TEXT,description TEXT,candidate_locations_json TEXT,status TEXT);
CREATE TABLE catalog_meta(key TEXT PRIMARY KEY,value TEXT);
CREATE INDEX variable_lookup ON variable_docs(release_id,table_name COLLATE NOCASE,variable_name COLLATE NOCASE);
CREATE INDEX code_lookup ON code_values(release_id,table_name COLLATE NOCASE,variable_name COLLATE NOCASE,code_value);
CREATE INDEX record_source ON records(source_id,sheet_id,source_row);
CREATE INDEX table_lookup ON table_docs(release_id,table_name COLLATE NOCASE);
'''
def add_record(db,release,source,sheet,row,kind,table,variable,title,body,raw):
    rid=db.execute('INSERT INTO records VALUES(NULL,?,?,?,?,?,?,?,?,?,?)',(release,source,sheet,row,kind,table,variable,title,body,js(raw))).lastrowid
    db.execute('INSERT INTO search_index VALUES(?,?,?,?,?,?,?)',(release,kind,table,variable,title,body,rid));return rid
def source(db,release,p,role):
    digest=sha(p)
    sid=db.execute('INSERT INTO source_files VALUES(NULL,?,?,?,?,?,0)',(release,role,str(p),p.stat().st_size,digest)).lastrowid
    return sid,digest
def role(name):
    s=name.lower()
    if s.startswith('vartable'):return 'variable_definition'
    if s.startswith('valuesets'):return 'code_value'
    if s.startswith('tables'):return 'table_definition'
    if 'introduction' in s:return 'introduction'
    if 'metadata' in s:return 'metadata_definition'
    if any(x in s for x in ['new','change']):return 'supplementary_change'
    return 'other_documentation'
def build():
    DATA.mkdir(parents=True,exist_ok=True)
    partial=TARGET.with_suffix('.sqlite.inprogress')
    if TARGET.exists() or partial.exists():raise RuntimeError('Catalog exists; preserve/version it before rebuilding: '+str(TARGET))
    db=sqlite3.connect(partial)
    try:
        db.executescript(DDL)
        for folder in sorted(ROOT.iterdir()):
            if not folder.is_dir() or not list(folder.glob('*.accdb')):continue
            release=folder.name;cycle=release.split('_')[1];label=release.rsplit('_',1)[-1].lower()
            db.execute('INSERT INTO releases VALUES(?,?,?,?,?)',(release,cycle,label,'Folder label, corroborated by paired ReadMe heading; table-level release details preserved separately',str(folder)))
            book=next(folder.glob('*.xlsx'));sid,digest=source(db,release,book,'codebook')
            workbook=openpyxl.load_workbook(book,read_only=True,data_only=False)
            for sheet in workbook:
                kind=role(sheet.title)
                sheetid=db.execute('INSERT INTO sheets VALUES(NULL,?,?,?,?,?)',(sid,sheet.title,kind,sheet.max_row,sheet.max_column)).lastrowid
                rows=sheet.iter_rows(values_only=True);header=None
                for n,row in enumerate(rows,1):
                    if not any(v is not None for v in row):continue
                    db.execute('INSERT INTO sheet_rows VALUES(?,?,?)',(sheetid,n,js(row)))
                    if n==1:header=[clean(v).lower() for v in row]
                    if kind in ['variable_definition','code_value','table_definition'] and n==1:continue
                    if kind in ['variable_definition','code_value','table_definition']:
                        raw={k:v for k,v in zip(header,row) if k};d={k:clean(v) for k,v in raw.items()}
                    else:
                        raw={'cells':row};d={}
                    table=d.get('tablename','');variable=d.get('varname','')
                    if kind=='table_definition':title=d.get('tabletitle',table)
                    elif kind=='variable_definition':title=d.get('vartitle',variable)
                    elif kind=='code_value':title=d.get('valuelabel','')
                    else:title=clean(next((v for v in row if v is not None),sheet.title))[:250]
                    body='\n'.join(k+': '+clean(v) for k,v in raw.items() if v is not None) if d else '\n'.join(clean(v) for v in row if v is not None)
                    rid=add_record(db,release,sid,sheetid,n,kind,table,variable,title,body,raw)
                    if kind=='table_definition' and table:
                        db.execute('INSERT INTO table_docs VALUES(?,?,?,?,?,?,?,?,?,?)',(rid,release,table,d.get('tablenumber'),d.get('survey'),title,d.get('description'),d.get('yearcoverage'),d.get('release'),d.get('release date')))
                    elif kind=='variable_definition' and table and variable:
                        db.execute('INSERT INTO variable_docs VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(rid,release,table,variable,d.get('varnumber'),title,d.get('longdescription'),d.get('datatype'),d.get('format'),d.get('fieldwidth'),d.get('multirecord'),d.get('imputationvar'),d.get('varsource'),d.get('filetitle'),d.get('sectiontitle')))
                    elif kind=='code_value' and table and variable:
                        db.execute('INSERT INTO code_values VALUES(?,?,?,?,?,?,?,?,?,?)',(rid,release,table,variable,d.get('varnumber'),d.get('codevalue'),title,d.get('frequency'),d.get('percent'),d.get('valueorder')))
            workbook.close()
            if sha(book)!=digest:raise RuntimeError('Codebook changed during import')
            db.execute('UPDATE source_files SET unchanged_after_read=1 WHERE source_id=?',(sid,))
            for doc in folder.glob('*.docx'):
                sourceid,digest=source(db,release,doc,'release_notes')
                with zipfile.ZipFile(doc) as z:
                    root=ET.fromstring(z.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                    paragraphs=[''.join(t.text or '' for t in p.findall('.//w:t',ns)) for p in root.findall('.//w:p',ns)]
                for n,line in enumerate(paragraphs,1):
                    if not line.strip():continue
                    rid=add_record(db,release,sourceid,None,n,'release_note','','',clean(line)[:250],line,{'paragraph':line})
                    db.execute('INSERT INTO release_notes VALUES(?,?,?,?)',(rid,release,n,line))
                if sha(doc)!=digest:raise RuntimeError('Release notes changed during import')
                db.execute('UPDATE source_files SET unchanged_after_read=1 WHERE source_id=?',(sourceid,))
            db.commit();print('CATALOG',release,flush=True)
        db.execute('INSERT INTO catalog_meta VALUES(?,?)',('scope','All original codebook sheets and release-note paragraphs; no harmonized measure mappings or analytical corrections applied.'))
        db.execute('INSERT INTO catalog_meta VALUES(?,?)',('builder_sha256',sha(Path(__file__))))
        db.commit()
        if db.execute('PRAGMA integrity_check').fetchone()[0]!='ok':raise RuntimeError('Catalog integrity failed')
        if db.execute('PRAGMA foreign_key_check').fetchall():raise RuntimeError('Catalog foreign keys failed')
    finally:db.close()
    partial.replace(TARGET)

def reconcile():
    db=sqlite3.connect(TARGET);db.execute('PRAGMA foreign_keys=ON')
    # Transaction makes reconciliation all-or-nothing. Original metadata is immutable.
    try:
        db.execute('DELETE FROM issues');db.execute('DELETE FROM physical_columns');db.execute('DELETE FROM physical_tables')
        for release, in db.execute('SELECT release_id FROM releases').fetchall():
            out=DATA/'annual'/release;manifest=json.loads((out/'verification.json').read_text())
            target=Path(manifest['output'])
            if sha(target)!=manifest['output_sha256']:raise RuntimeError('Converted file hash mismatch '+release)
            source_path=Path(manifest['source'])
            if sha(source_path)!=manifest['source_sha256']:raise RuntimeError('Original source changed '+release)
            existing=db.execute("SELECT source_id FROM source_files WHERE release_id=? AND role='access_database'",(release,)).fetchone()
            if existing:
                db.execute('UPDATE source_files SET path=?,bytes=?,sha256=?,unchanged_after_read=1 WHERE source_id=?',(str(source_path),manifest['source_bytes'],manifest['source_sha256'],existing[0]))
            else:db.execute('INSERT INTO source_files VALUES(NULL,?,?,?,?,?,1)',(release,'access_database',str(source_path),manifest['source_bytes'],manifest['source_sha256']))
            checks=list(csv.DictReader((out/'table_checks.csv').open(encoding='utf-8',newline='')))
            for t in checks:db.execute('INSERT INTO physical_tables VALUES(?,?,?,?,?,?,?)',(release,t['table'],int(t['source_rows']),int(t['sqlite_rows']),str(target),1,int(t['source_values_sha256']==t['sqlite_values_sha256'])))
            with (out/'columns.csv').open(encoding='utf-8',newline='') as f:
                for col in csv.DictReader(f):db.execute('INSERT INTO physical_columns VALUES(?,?,?,?,?,?,?,?)',(release,col['table'],col['column'],int(col['ordinal']),col['access_type'],int(col['access_width']),col['sqlite_type'],int(col['null_count'])))
        def issue(release,kind,severity,rid,table,variable,desc,candidates):
            db.execute('INSERT INTO issues VALUES(NULL,?,?,?,?,?,?,?,?,?)',(release,kind,severity,rid,table,variable,desc,js(candidates),'observed_not_corrected'))
        for rid,release,table in db.execute('SELECT record_id,release_id,table_name FROM table_docs').fetchall():
            if not db.execute('SELECT 1 FROM physical_tables WHERE release_id=? AND table_name=? COLLATE NOCASE',(release,table)).fetchone():
                issue(release,'documented_table_absent','high',rid,table,'','Codebook table name has no physical table in the supplied Access source.',[])
        for rid,release,table,variable in db.execute('SELECT record_id,release_id,table_name,variable_name FROM variable_docs').fetchall():
            physical=db.execute('SELECT 1 FROM physical_tables WHERE release_id=? AND table_name=? COLLATE NOCASE',(release,table)).fetchone()
            if physical and not db.execute('SELECT 1 FROM physical_columns WHERE release_id=? AND table_name=? COLLATE NOCASE AND column_name=? COLLATE NOCASE',(release,table,variable)).fetchone():
                candidates=[r[0] for r in db.execute('SELECT table_name FROM physical_columns WHERE release_id=? AND column_name=? COLLATE NOCASE',(release,variable))]
                issue(release,'documented_column_absent','high',rid,table,variable,'Documented placement not found. Candidate locations are name matches only, not approved semantic corrections.',candidates)
        for release,table,column in db.execute('SELECT release_id,table_name,column_name FROM physical_columns').fetchall():
            if not db.execute('SELECT 1 FROM table_docs WHERE release_id=? AND table_name=? COLLATE NOCASE',(release,table)).fetchone():continue
            if not db.execute('SELECT 1 FROM variable_docs WHERE release_id=? AND table_name=? COLLATE NOCASE AND variable_name=? COLLATE NOCASE',(release,table,column)).fetchone():
                issue(release,'physical_column_not_in_dictionary','informational',None,table,column,'Physical column lacks a variable entry at this table location. Shared identifiers and helper fields may legitimately be omitted.',[])
        db.execute("INSERT OR REPLACE INTO catalog_meta VALUES('physical_reconciliation','Complete for all imported releases; issues retained separately; no automatic source corrections')")
        db.commit()
    except Exception:db.rollback();raise
    finally:db.close()

def verify():
    db=sqlite3.connect('file:'+TARGET.as_posix()+'?mode=ro',uri=True)
    report={'catalog_path':str(TARGET),'counts':{},'integrity_check':db.execute('PRAGMA integrity_check').fetchone()[0],'foreign_key_violations':len(db.execute('PRAGMA foreign_key_check').fetchall())}
    for table in ['releases','source_files','sheets','sheet_rows','records','table_docs','variable_docs','code_values','release_notes','physical_tables','physical_columns','issues']:
        report['counts'][table]=db.execute('SELECT COUNT(*) FROM '+table).fetchone()[0]
    report['record_search_count_matches']=db.execute('SELECT COUNT(*) FROM records').fetchone()[0]==db.execute('SELECT COUNT(*) FROM search_index').fetchone()[0]
    report['all_document_sources_unchanged']=db.execute('SELECT COUNT(*) FROM source_files WHERE unchanged_after_read<>1').fetchone()[0]==0
    report['keyword_search_smoke_checks']={term:db.execute('SELECT COUNT(*) FROM search_index WHERE search_index MATCH ?',(term,)).fetchone()[0] for term in ['financial aid','graduation','CIPCODE','tuition']}
    report['issues_by_type']=dict(db.execute('SELECT issue_type,COUNT(*) FROM issues GROUP BY issue_type'))
    report['annual_coverage']=[dict(zip(['release','sheets','variable_definitions','categorical_codes'],r)) for r in db.execute('SELECT r.release_id,(SELECT COUNT(*) FROM sheets s JOIN source_files f ON s.source_id=f.source_id WHERE f.release_id=r.release_id),(SELECT COUNT(*) FROM variable_docs v WHERE v.release_id=r.release_id),(SELECT COUNT(*) FROM code_values c WHERE c.release_id=r.release_id) FROM releases r ORDER BY release_id')]
    assert report['integrity_check']=='ok' and report['foreign_key_violations']==0 and report['record_search_count_matches'] and report['all_document_sources_unchanged']
    assert report['counts']['releases']==10 and all(n>0 for n in report['keyword_search_smoke_checks'].values())
    (DATA/'catalog_verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8');db.close();print(json.dumps(report,indent=2))
if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--reconcile',action='store_true');parser.add_argument('--verify-only',action='store_true');args=parser.parse_args()
    if args.reconcile:reconcile()
    elif not args.verify_only:build()
    verify()
